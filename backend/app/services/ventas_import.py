"""Lectura del extracto de ventas (ISVAN / TRALOG) para el modulo de
indicadores.

Acepta los archivos tal como los arma el negocio, sin pedir una plantilla:

- El mensual: hoja "data" con el encabezado en la fila 1, mas la hoja
  "precio de compras" con el costo por caja (encabezado en la fila 10,
  debajo del membrete).
- El libro de cierre: la misma hoja "data" pero con totales encima del
  encabezado, y decenas de hojas de tablas dinamicas alrededor.
- El extracto diario: una sola hoja con los encabezados en minuscula
  ("num docum", "divisas", ...) y a veces el costo en una columna "pc".
- El reporte tal como sale del sistema: sin encabezado y en otro orden. Las
  columnas se reconocen por su contenido (ver reconocer_columnas).

Este modulo solo lee y normaliza. Lo que depende de la base (clientes
nuevos, solapes con otras cargas, costos ya cargados) lo resuelve
app/api/indicadores.py.
"""

from __future__ import annotations

import io
import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass
from functools import cached_property
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter
from fastapi import HTTPException

from app.core.excel_utils import mapear_columnas, normalizar_encabezado, valor_a_texto

# Los alias ya vienen normalizados (minusculas, sin acentos, espacios -> _).
# Cubren los tres formatos, incluidas las erratas reales de los archivos
# ("provedor", "tipo de cloente").
ALIAS_COLUMNAS: dict[str, list[str]] = {
    "ruta": ["ruta", "numero_ruta"],
    "codigo_cliente": ["cod_cte", "codigo_cliente", "cod_cliente"],
    "cliente": ["cliente", "nombre_cliente"],
    "codigo_producto": ["cod_producto", "codigo_producto"],
    "producto": ["producto", "nombre_producto"],
    "grupo": ["grupo"],
    "num_doc": ["num_doc", "num_docum", "numero_documento"],
    "fecha": ["fecha"],
    "cajas": ["cajas"],
    "unidades": ["und", "unidades"],
    "litros": ["litros"],
    "monto_bs": ["bolivares", "monto_bs"],
    "monto_usd": ["divisas", "monto_usd"],
    "tipo_doc": ["tp_doc", "tipo_de_documento", "cod_tipo_doc"],
    "tipo_cliente": ["tipo_de_cliente", "tipo_de_cloente", "tipo_cliente"],
    "lista_precio": ["precio", "lista_precio"],
    "proveedor": ["proveedor", "provedor", "codigo_proveedor"],
    "costo_caja": ["pc", "costo_caja_usd"],
}
CAMPOS_REQUERIDOS = [
    "ruta", "codigo_cliente", "cliente", "codigo_producto", "producto", "grupo",
    "num_doc", "fecha", "cajas", "unidades", "litros", "monto_bs", "monto_usd",
    "tipo_doc", "tipo_cliente",
]

# FA (factura) y NE (nota de entrega) son ventas; DV y DN, devoluciones con
# montos negativos. Cualquier otro codigo es un dato que no sabemos leer.
TIPO_MOVIMIENTO_POR_DOC = {"FA": "VENTA", "NE": "VENTA", "DV": "DEVOLUCION", "DN": "DEVOLUCION"}

# Hasta donde se busca la fila de encabezado: el libro de cierre trae
# totales arriba y la hoja de costos, un membrete de 9 filas.
FILAS_ANTES_DEL_ENCABEZADO = 15


@dataclass
class FilaVenta:
    fila: int
    fecha: date
    num_doc: str
    cod_tipo_doc: str
    tipo_movimiento: str
    codigo_cliente: str
    codigo_producto: int
    cajas: float
    unidades: int
    litros: float
    monto_bs: float
    monto_usd: float


@dataclass
class ExtractoVentas:
    hoja: str
    # Otras hojas del libro con el mismo formato que no se leyeron: si el
    # usuario esperaba que se sumaran, tiene que saberlo antes de confirmar.
    otras_hojas_con_formato: list[str]
    filas: list[FilaVenta]
    errores: list[dict]
    # codigo -> {nombre, tipo, ruta, listaPrecio}
    clientes: dict[str, dict]
    # codigo -> {nombre, grupo, codigoProveedor, unidadesPorCaja, litrosPorUnidad, costoCajaUsd}
    productos: dict[int, dict]
    # De donde salio el costo por caja (None: el archivo no lo trae).
    fuente_costo: str | None
    # False si el archivo vino sin encabezado y las columnas se reconocieron
    # por su contenido.
    con_encabezado: bool
    # Que columna se tomo para cada dato (ver describir_columnas).
    columnas: list[dict]


def _texto(valor: object) -> str:
    """Quita los espacios de relleno: los nombres vienen completados con
    blancos hasta 120 caracteres y los encabezados con espacios al final."""
    return " ".join(str(valor).split()) if valor is not None else ""


def _a_numero(valor: object) -> float | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        return float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None


def _a_fecha(valor: object) -> date | None:
    # En el archivo mensual la fecha viene como TEXTO ("2026-06-12"); en
    # otros extractos, como fecha de Excel.
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    return None


def _codigo_producto(valor: object) -> int | None:
    try:
        return int(valor_a_texto(valor))
    except ValueError:
        return None


def _codigo_cliente(valor: object) -> str:
    # "000758" en el reporte del sistema, 758 en los extractos con
    # encabezado: sin ceros a la izquierda quedan iguales (y cruzan con el
    # maestro de clientes de logistica).
    texto = valor_a_texto(valor)
    return texto.lstrip("0") or ("0" if texto else "")


def _num_doc(valor: object) -> str:
    # "000001237" en el mensual, 1237 en el diario: sin ceros a la izquierda
    # los dos formatos quedan iguales (y como en Despacho.numeroDocumento).
    texto = valor_a_texto(valor)
    return texto.lstrip("0") or ("0" if texto else "")


def _es_encabezado_de_ventas(fila: tuple) -> bool:
    nombres = {normalizar_encabezado(v) for v in fila if v is not None}
    return "fecha" in nombres and bool(nombres & {"divisas", "monto_usd"})


def _buscar_encabezado(hoja, es_encabezado) -> tuple[int, tuple] | None:
    for i, fila in enumerate(hoja.iter_rows(values_only=True, max_row=FILAS_ANTES_DEL_ENCABEZADO)):
        if fila and es_encabezado(fila):
            return i, fila
    return None


def _hojas_por_prioridad(libro) -> list:
    """La hoja "data" primero (asi se llama en el mensual y en el cierre);
    despues el resto, en el orden del libro."""
    return sorted(libro.worksheets, key=lambda h: normalizar_encabezado(h.title) != "data")


def _leer_costos(libro, hoja_ventas) -> dict[int, float]:
    """Costo por CAJA desde la hoja "precio de compras" (columnas codigo y
    "Pc"). Vacio si el libro no la trae."""

    def es_encabezado(fila: tuple) -> bool:
        nombres = {normalizar_encabezado(v) for v in fila if v is not None}
        return "pc" in nombres and any(n.startswith("codigo") for n in nombres)

    candidatas = sorted(
        (h for h in libro.worksheets if h.title != hoja_ventas.title),
        key=lambda h: not normalizar_encabezado(h.title).startswith("precio_de_compra"),
    )
    for hoja in candidatas:
        encontrado = _buscar_encabezado(hoja, es_encabezado)
        if not encontrado:
            continue
        indice, encabezado = encontrado
        nombres = [normalizar_encabezado(v) for v in encabezado]
        col_codigo = next(i for i, n in enumerate(nombres) if n.startswith("codigo"))
        col_costo = nombres.index("pc")
        costos: dict[int, float] = {}
        for fila in hoja.iter_rows(values_only=True, min_row=indice + 2):
            if not fila or len(fila) <= max(col_codigo, col_costo):
                continue
            codigo = _codigo_producto(fila[col_codigo])
            costo = _a_numero(fila[col_costo])
            if codigo is not None and costo is not None and costo > 0:
                costos[codigo] = costo
        if costos:
            return costos
    return {}


# ---------- Archivos sin encabezado ----------
#
# El reporte que el cliente descarga de su sistema viene sin fila de
# encabezado y con las columnas en otro orden. En vez de exigirles una
# plantilla, cada columna se reconoce por lo que contiene. Las pruebas son
# relaciones entre columnas, no formatos sueltos: "cada codigo tiene un solo
# nombre", "cada documento es de un solo cliente", "bolivares / divisas da
# una tasa estable", "unidades / cajas es siempre entero"... Asi no importa
# el orden, y dos columnas parecidas (el codigo de cliente y el numero de
# documento son los dos numeros con ceros) no se confunden.
#
# Si una columna obligatoria no aparece, el archivo se rechaza diciendo cual
# falta: nunca se completa con un valor inventado.

# Nombre de cada dato tal como se le muestra al usuario.
NOMBRE_CAMPO: dict[str, str] = {
    "ruta": "ruta",
    "codigo_cliente": "codigo de cliente",
    "cliente": "nombre del cliente",
    "codigo_producto": "codigo de producto (SKU)",
    "producto": "nombre del producto",
    "grupo": "grupo de producto",
    "num_doc": "numero de documento",
    "fecha": "fecha",
    "cajas": "cajas",
    "unidades": "unidades",
    "litros": "litros",
    "monto_bs": "monto en bolivares",
    "monto_usd": "monto en divisas (USD)",
    "tipo_doc": "tipo de documento (FA, NE, DV, DN)",
    "tipo_cliente": "tipo de cliente",
    "lista_precio": "lista de precio",
    "proveedor": "proveedor",
    "costo_caja": "costo por caja",
}

# Cuantas filas se miran para reconocer las columnas: sobra para que las
# relaciones se vean, y no depende del tamano del archivo.
FILAS_PARA_RECONOCER = 3000
# Proporcion de filas que tiene que cumplir una prueba: tolera algun dato
# suelto mal cargado sin dejar de reconocer la columna.
CASI_TODAS = 0.95

_PRESENTACION = re.compile(r"\d\s*X\s*\d|\d\s*(ML|LTS?|L|GRS?|KGS?)\b", re.IGNORECASE)
_RUTA = re.compile(r"R?\d{1,3}", re.IGNORECASE)
_DIGITOS = re.compile(r"\d+")
_LETRA = re.compile(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]")
# Filtro barato antes de intentar leer la fecha de verdad (strptime es lento
# y se prueba en todas las columnas).
_PARECE_FECHA = re.compile(r"\d{1,4}[-/]\d{1,2}[-/]\d{1,4}")


def _es_numero(valor: object) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def _es_entero(x: float) -> bool:
    return abs(x - round(x)) < 1e-6


class _Columna:
    """Una columna de la muestra, con lo que se le pregunta una y otra vez
    ya calculado: las pruebas comparan cada columna contra todas las demas,
    y recalcularlo en cada comparacion hacia la lectura 20 veces mas lenta."""

    def __init__(self, indice: int, valores: list):
        self.indice = indice
        self.valores = valores
        self.textos = [_texto(valor_a_texto(v)) if v is not None else "" for v in valores]
        self.numeros = [float(v) if _es_numero(v) else None for v in valores]
        no_vacios = [t for t in self.textos if t]
        self.distintos = len(set(no_vacios))
        self.vacia = not no_vacios
        self._casos = [(v, t) for v, t in zip(valores, self.textos) if t]

    def proporcion(self, prueba) -> float:
        return sum(1 for v, t in self._casos if prueba(v, t)) / len(self._casos) if self._casos else 0.0

    def numero(self, i: int) -> float | None:
        return self.numeros[i]

    @cached_property
    def son_digitos(self) -> bool:
        return self.proporcion(lambda v, t: bool(_DIGITOS.fullmatch(t))) >= CASI_TODAS

    @cached_property
    def son_textos(self) -> bool:
        return self.proporcion(lambda v, t: bool(_LETRA.search(t))) >= 0.8

    @cached_property
    def son_numeros(self) -> bool:
        return self.distintos > 1 and self.proporcion(lambda v, t: _es_numero(v)) >= CASI_TODAS

    @cached_property
    def es_nombre_de_producto(self) -> bool:
        return self.proporcion(lambda v, t: bool(_PRESENTACION.search(t))) >= 0.5

    @cached_property
    def empieza_con_digito(self) -> float:
        """Para no tomar como nombre una columna de "codigo + nombre" pegados
        ("67888084TIO RICO MANTECADO"), que algunos extractos traen aparte."""
        return self.proporcion(lambda v, t: t[:1].isdigit())


def _determina(a: _Columna, b: _Columna) -> float:
    """Que tanto el valor de `a` fija el de `b`: la fraccion de valores de
    `a` que aparecen siempre junto al mismo valor de `b`."""
    pares: dict[str, set] = defaultdict(set)
    for x, y in zip(a.textos, b.textos):
        if x and y:
            pares[x].add(y)
    return sum(1 for s in pares.values() if len(s) == 1) / len(pares) if pares else 0.0


def reconocer_columnas(filas: list[tuple]) -> dict[str, int]:
    """Campo -> indice de columna, para una hoja sin encabezado. Lanza 400
    con la lista de lo que no se pudo reconocer si falta algo obligatorio."""
    muestra = [f for f in filas if f and any(v is not None and _texto(v) for v in f)][:FILAS_PARA_RECONOCER]
    if not muestra:
        raise HTTPException(400, "La hoja de ventas esta vacia")
    total_filas = len(muestra)
    ancho = max(len(f) for f in muestra)
    columnas = [_Columna(i, [f[i] if i < len(f) else None for f in muestra]) for i in range(ancho)]
    libres = [c for c in columnas if not c.vacia]
    mapa: dict[str, int] = {}

    def asignar(campo: str, columna: _Columna) -> None:
        mapa[campo] = columna.indice
        libres.remove(columna)

    def columna_de(campo: str) -> _Columna | None:
        return columnas[mapa[campo]] if campo in mapa else None

    # Tipo de documento y fecha: se reconocen por su propio valor.
    for c in list(libres):
        if c.proporcion(lambda v, t: t.upper() in TIPO_MOVIMIENTO_POR_DOC) >= CASI_TODAS:
            asignar("tipo_doc", c)
            break
    for c in list(libres):
        es_fecha = lambda v, t: isinstance(v, (date, datetime)) or (  # noqa: E731
            isinstance(v, str) and bool(_PARECE_FECHA.match(t)) and _a_fecha(v) is not None
        )
        if c.proporcion(es_fecha) >= CASI_TODAS:
            asignar("fecha", c)
            break

    # La semana (numero ISO de la fecha) no se usa: se calcula de la fecha.
    # Se aparta para que no se confunda con la ruta, que tambien son numeros
    # de dos cifras.
    if "fecha" in mapa:
        fechas = [_a_fecha(v) for v in columna_de("fecha").valores]
        for c in list(libres):
            coinciden = [
                f is not None and t.isdigit() and int(t) == f.isocalendar()[1]
                for f, t in zip(fechas, c.textos)
                if t
            ]
            if coinciden and sum(coinciden) / len(coinciden) >= 0.9:
                libres.remove(c)
                break

    # Codigo + nombre, de producto y de cliente: cada codigo tiene un solo
    # nombre y casi cada nombre un solo codigo (una cadena repite nombre con
    # varios codigos de sucursal). El producto es el que tiene presentacion
    # en el nombre ("18X90ML", "1X3.6L").
    pares = [
        (d, n)
        for d in libres
        if d.son_digitos and d.distintos >= 2
        for n in libres
        if n is not d
        and n.son_textos
        and n.distintos >= max(2, 0.5 * d.distintos)
        and _determina(d, n) >= 0.98
        and _determina(n, d) >= 0.8
    ]
    # Entre dos nombres posibles para el mismo codigo gana el que no empieza
    # con numeros (el nombre limpio, no "codigo + nombre").
    preferencia = lambda par: (par[0].distintos, -par[1].empieza_con_digito, -par[1].indice)  # noqa: E731
    de_producto = [par for par in pares if par[1].es_nombre_de_producto]
    if de_producto:
        codigo, nombre = max(de_producto, key=preferencia)
        asignar("codigo_producto", codigo)
        asignar("producto", nombre)
    de_cliente = [
        par for par in pares if not par[1].es_nombre_de_producto and par[0] in libres and par[1] in libres
    ]
    if de_cliente:
        codigo, nombre = max(de_cliente, key=preferencia)
        asignar("codigo_cliente", codigo)
        asignar("cliente", nombre)

    # Bolivares y divisas: su cociente es la tasa del dia, estable en casi
    # todas las filas (con alguna fila en tasa 1 que la validacion reporta).
    numericas = [c for c in libres if c.son_numeros]
    mejor_monto = None
    for bs in numericas:
        for usd in numericas:
            if bs is usd:
                continue
            tasas = [
                bs.numero(i) / usd.numero(i)
                for i in range(total_filas)
                if bs.numero(i) is not None and usd.numero(i)
            ]
            if len(tasas) < 0.5 * total_filas:
                continue
            mediana = statistics.median(tasas)
            if mediana < 5:
                continue
            estables = sum(1 for t in tasas if 0.5 * mediana <= t <= 1.5 * mediana) / len(tasas)
            if estables >= 0.9 and (mejor_monto is None or estables > mejor_monto[0]):
                mejor_monto = (estables, bs, usd)
    if mejor_monto:
        asignar("monto_bs", mejor_monto[1])
        asignar("monto_usd", mejor_monto[2])

    # Unidades y cajas: unidades es entero y unidades / cajas tambien (las
    # unidades de la caja), con el mismo signo. Si hay mas de un par posible,
    # las cajas son la columna de menor volumen.
    numericas = [c for c in libres if c.son_numeros]
    pares_cajas = []
    for unidades in numericas:
        if unidades.proporcion(lambda v, t: _es_numero(v) and _es_entero(float(v))) < 0.99:
            continue
        for cajas in numericas:
            if cajas is unidades:
                continue
            juntos = [
                (unidades.numero(i), cajas.numero(i))
                for i in range(total_filas)
                if unidades.numero(i) is not None and cajas.numero(i)
            ]
            if len(juntos) < 0.5 * total_filas:
                continue
            cumplen = sum(1 for u, k in juntos if _es_entero(u / k) and u / k > 0 and abs(k) <= abs(u) + 1e-9)
            if cumplen / len(juntos) >= 0.98:
                pares_cajas.append((sum(abs(k) for _, k in juntos), unidades, cajas))
    if pares_cajas:
        _, unidades, cajas = min(pares_cajas, key=lambda par: par[0])
        asignar("unidades", unidades)
        asignar("cajas", cajas)

    # Litros: por unidad dan siempre lo mismo para un mismo producto (el
    # tamano de la presentacion). Descarta columnas parecidas, como una de
    # mililitros cargada a veces en litros y a veces no.
    if "unidades" in mapa and "codigo_producto" in mapa:
        unidades, producto = columna_de("unidades"), columna_de("codigo_producto")
        mejor_litros = None
        for litros in [c for c in libres if c.son_numeros]:
            por_producto: dict[str, set] = defaultdict(set)
            razonables = casos = 0
            for i in range(total_filas):
                u, l = unidades.numero(i), litros.numero(i)
                if not u or l is None:
                    continue
                casos += 1
                razonables += 0 <= l / u <= 25
                por_producto[producto.textos[i]].add(round(l / u, 3))
            if not casos or razonables / casos < 0.9:
                continue
            constancia = sum(1 for valores in por_producto.values() if len(valores) == 1) / len(por_producto)
            if constancia >= 0.8 and (mejor_litros is None or constancia > mejor_litros[0]):
                mejor_litros = (constancia, litros)
        if mejor_litros:
            asignar("litros", mejor_litros[1])

    # Numero de documento: cada documento es de un solo cliente y tiene
    # varias lineas (los montos, que tambien serian "unicos", no se repiten).
    # Cada tipo lleva su propia numeracion (la factura 100 y la nota de
    # entrega 100 son de clientes distintos), asi que el documento es el par
    # tipo + numero.
    cliente = columna_de("codigo_cliente")
    tipo_doc = columna_de("tipo_doc")

    def con_tipo(c: _Columna) -> _Columna:
        if tipo_doc is None:
            return c
        return _Columna(c.indice, [f"{t}|{n}" if n else None for t, n in zip(tipo_doc.textos, c.textos)])

    if cliente:
        documentos = [
            c
            for c in libres
            if c.son_digitos
            and c.distintos >= 2
            and c.distintos < total_filas
            and _determina(con_tipo(c), cliente) >= 0.98
        ]
        if documentos:
            asignar("num_doc", max(documentos, key=lambda c: c.distintos))

    # Ruta: codigos cortos (R1, 11, 77) que casi no cambian por cliente.
    rutas = [
        c
        for c in libres
        if c.distintos <= 100
        and c.proporcion(lambda v, t: bool(_RUTA.fullmatch(t))) >= CASI_TODAS
        and (cliente is None or _determina(cliente, c) >= 0.9)
    ]
    if rutas:
        asignar("ruta", max(rutas, key=lambda c: c.distintos))

    # Lista de precio: una letra (A, B, C).
    for c in list(libres):
        if c.distintos <= 10 and c.proporcion(lambda v, t: bool(re.fullmatch(r"[A-Za-z]", t))) >= CASI_TODAS:
            asignar("lista_precio", c)
            break

    # Grupo de producto y tipo de cliente: pocos valores de texto. El grupo
    # lo fija el producto; el tipo, el cliente (a un cliente se le venden
    # varios grupos y un producto se le vende a varios tipos).
    producto = columna_de("codigo_producto")
    categorias = [c for c in libres if c.son_textos and c.distintos <= 60]
    if producto:
        # Un grupo junta varios productos: una columna que va 1 a 1 con el
        # producto es otro nombre del producto (p. ej. "codigo + nombre").
        grupos = [
            c
            for c in categorias
            if _determina(producto, c) >= 0.98 and (producto.distintos <= 1 or _determina(c, producto) < 0.98)
        ]
        if cliente:
            grupos = [c for c in grupos if _determina(producto, c) > _determina(cliente, c) or c.distintos == 1]
        if grupos:
            asignar("grupo", max(grupos, key=lambda c: c.distintos))
    if cliente:
        tipos = [
            c
            for c in libres
            if c.son_textos
            and c.distintos <= 20
            and c.proporcion(lambda v, t: len(t) >= 3) >= CASI_TODAS
            and _determina(cliente, c) >= 0.98
            and (producto is None or _determina(producto, c) < 0.98 or c.distintos == 1)
        ]
        if tipos:
            asignar("tipo_cliente", max(tipos, key=lambda c: c.distintos))

    # Proveedor: codigo numerico fijo por producto, con pocos valores.
    if producto:
        for c in list(libres):
            if c.son_digitos and c.distintos <= 50 and _determina(producto, c) >= 0.98:
                asignar("proveedor", c)
                break

    faltantes = [campo for campo in CAMPOS_REQUERIDOS if campo not in mapa]
    if faltantes:
        nombres = [NOMBRE_CAMPO[c] for c in faltantes]
        # Detalle estructurado: la pantalla muestra primero lo que falta y deja
        # a mano (desplegable) lo que si se reconocio. Como texto corrido, la
        # lista de 15 columnas reconocidas tapaba el unico dato importante.
        raise HTTPException(
            400,
            {
                "mensaje": (
                    f"Falta {'la columna obligatoria' if len(nombres) == 1 else 'las columnas obligatorias'}: "
                    f"{', '.join(nombres)}. Agrégala al archivo (en cualquier posición) y vuelve a subirlo."
                ),
                "columnasFaltantes": nombres,
                "columnasReconocidas": describir_columnas(muestra, mapa, None),
            },
        )
    return mapa


def describir_columnas(filas: list[tuple], mapa: dict[str, int], encabezado: tuple | None) -> list[dict]:
    """Que columna del archivo se tomo para cada dato, con ejemplos, para
    que el ADMIN lo compruebe de un vistazo antes de confirmar."""
    descripcion = []
    for campo, indice in sorted(mapa.items(), key=lambda item: item[1]):
        ejemplos: list[str] = []
        for fila in filas:
            if indice < len(fila) and fila[indice] is not None:
                texto = _texto(valor_a_texto(fila[indice]))
                if texto and texto not in ejemplos:
                    ejemplos.append(texto[:40])
            if len(ejemplos) == 3:
                break
        titulo = _texto(encabezado[indice]) if encabezado and indice < len(encabezado) else None
        descripcion.append(
            {
                "campo": NOMBRE_CAMPO[campo],
                "columna": get_column_letter(indice + 1),
                "encabezado": titulo or None,
                "ejemplos": ejemplos,
            }
        )
    return descripcion


def leer_extracto(contenido: bytes) -> ExtractoVentas:
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo — verifica que sea un Excel (.xlsx) valido")

    try:
        return _leer_libro(libro)
    finally:
        libro.close()


def _hoja_sin_encabezado(libro):
    """Ninguna hoja trae encabezado: se reconoce la hoja de ventas por su
    contenido. Primero la hoja "data" (asi se llama tambien en el reporte
    del sistema), despues el resto; nunca la de costos."""
    candidatas = [
        h for h in _hojas_por_prioridad(libro)
        if not normalizar_encabezado(h.title).startswith("precio_de_compra")
    ]
    if not candidatas:
        raise HTTPException(400, "El libro no tiene una hoja de ventas")
    primer_error: HTTPException | None = None
    for hoja in candidatas:
        filas = list(hoja.iter_rows(values_only=True))
        try:
            mapa = reconocer_columnas(filas)
        except HTTPException as error:
            primer_error = primer_error or error
            continue
        return hoja, mapa, filas, []
    # El motivo que se muestra es el de la hoja mas probable (la primera).
    raise primer_error


def _leer_libro(libro) -> ExtractoVentas:
    hoja = None
    encabezado: tuple = ()
    indice_encabezado = 0
    otras: list[str] = []
    for candidata in _hojas_por_prioridad(libro):
        encontrado = _buscar_encabezado(candidata, _es_encabezado_de_ventas)
        if not encontrado:
            continue
        if hoja is None:
            hoja = candidata
            indice_encabezado, encabezado = encontrado
        else:
            otras.append(candidata.title)
    if hoja is not None:
        mapa = mapear_columnas(encabezado, ALIAS_COLUMNAS, CAMPOS_REQUERIDOS)
        filas_hoja = list(hoja.iter_rows(values_only=True, min_row=indice_encabezado + 2))
        primera_fila_de_datos = indice_encabezado + 2  # 1-based, como la ve el usuario en Excel
    else:
        hoja, mapa, filas_hoja, otras = _hoja_sin_encabezado(libro)
        encabezado = None
        primera_fila_de_datos = 1
    columnas = describir_columnas(filas_hoja, mapa, encabezado)
    costos_hoja = _leer_costos(libro, hoja)

    filas: list[FilaVenta] = []
    errores: list[dict] = []
    clientes: dict[str, dict] = {}
    fecha_cliente: dict[str, date] = {}
    productos: dict[int, dict] = {}
    unidades_por_caja: dict[int, Counter] = {}
    litros_por_unidad: dict[int, Counter] = {}
    costo_columna: dict[int, float] = {}

    for n, fila in enumerate(filas_hoja, start=primera_fila_de_datos):
        if not fila or all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
            continue

        def val(campo: str):
            idx = mapa.get(campo)
            return fila[idx] if idx is not None and idx < len(fila) else None

        codigo_cliente = _codigo_cliente(val("codigo_cliente"))
        codigo_producto = _codigo_producto(val("codigo_producto"))
        num_doc = _num_doc(val("num_doc"))
        fecha = _a_fecha(val("fecha"))
        cod_tipo_doc = _texto(val("tipo_doc")).upper()
        tipo_movimiento = TIPO_MOVIMIENTO_POR_DOC.get(cod_tipo_doc)
        monto_usd = _a_numero(val("monto_usd"))
        cajas = _a_numero(val("cajas"))
        unidades = _a_numero(val("unidades"))

        error: tuple[str, str] | None = None
        if not codigo_cliente:
            error = ("codigo cliente", "Falta el codigo de cliente")
        elif codigo_producto is None:
            error = ("codigo producto", f'Codigo de producto invalido "{_texto(val("codigo_producto"))}"')
        elif not num_doc:
            error = ("num doc", "Falta el numero de documento")
        elif fecha is None:
            error = ("fecha", f'Fecha invalida "{_texto(val("fecha"))}"')
        elif tipo_movimiento is None:
            error = ("tipo de documento", f'Tipo de documento desconocido "{cod_tipo_doc}" (se esperan FA, NE, DV o DN)')
        elif monto_usd is None:
            error = ("divisas", "Falta el monto en divisas")
        elif cajas is None or unidades is None:
            error = ("cajas/unidades", "Faltan las cajas o las unidades")
        if error:
            errores.append({"fila": n, "columna": error[0], "motivo": error[1]})
            continue

        litros = _a_numero(val("litros")) or 0.0
        filas.append(
            FilaVenta(
                fila=n,
                fecha=fecha,
                num_doc=num_doc,
                cod_tipo_doc=cod_tipo_doc,
                tipo_movimiento=tipo_movimiento,
                codigo_cliente=codigo_cliente,
                codigo_producto=codigo_producto,
                cajas=cajas,
                unidades=int(round(unidades)),
                litros=litros,
                monto_bs=_a_numero(val("monto_bs")) or 0.0,
                monto_usd=monto_usd,
            )
        )

        # El cliente queda con los datos de su fila mas reciente: si en el
        # periodo le cambiaron la ruta, manda la ultima.
        if codigo_cliente not in fecha_cliente or fecha >= fecha_cliente[codigo_cliente]:
            fecha_cliente[codigo_cliente] = fecha
            clientes[codigo_cliente] = {
                "nombre": _texto(val("cliente")),
                "tipo": _texto(val("tipo_cliente")).upper(),
                "ruta": _texto(val("ruta")).upper(),
                "listaPrecio": _texto(val("lista_precio")) or None,
            }

        productos.setdefault(
            codigo_producto,
            {
                "nombre": _texto(val("producto")),
                "grupo": _texto(val("grupo")).upper(),
                "codigoProveedor": valor_a_texto(val("proveedor")) or None,
            },
        )
        if cajas and unidades:
            unidades_por_caja.setdefault(codigo_producto, Counter())[round(abs(unidades / cajas))] += 1
        if unidades and litros:
            litros_por_unidad.setdefault(codigo_producto, Counter())[round(abs(litros / unidades), 4)] += 1
        costo = _a_numero(val("costo_caja"))
        if costo:
            costo_columna[codigo_producto] = costo

    if not filas and not errores:
        raise HTTPException(400, f'La hoja "{hoja.title}" no tiene filas de ventas')

    if costos_hoja:
        fuente_costo = 'hoja "precio de compras"'
        costos = costos_hoja
    elif costo_columna:
        fuente_costo = 'columna "pc"'
        costos = costo_columna
    else:
        fuente_costo = None
        costos = {}

    for codigo, producto in productos.items():
        producto["unidadesPorCaja"] = (
            unidades_por_caja[codigo].most_common(1)[0][0] if codigo in unidades_por_caja else None
        )
        producto["litrosPorUnidad"] = (
            litros_por_unidad[codigo].most_common(1)[0][0] if codigo in litros_por_unidad else None
        )
        producto["costoCajaUsd"] = costos.get(codigo)

    return ExtractoVentas(
        hoja=hoja.title,
        otras_hojas_con_formato=otras,
        filas=filas,
        errores=errores,
        clientes=clientes,
        productos=productos,
        fuente_costo=fuente_costo,
        con_encabezado=encabezado is not None,
        columnas=columnas,
    )
