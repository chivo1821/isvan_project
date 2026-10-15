"""Importacion del historico del censo hecho en Survey123 (la plataforma
de formularios que se usaba antes de la pagina).

La exportacion trae una fila por censo con los nombres de las preguntas del
formulario como columnas (los mismos de app/formularios/censo.json), mas
smid (el id de la fila alla, estable entre exportaciones), x/y y ubicacion
(WKT). Se acepta en CSV (con ; o ,), GeoJSON o Excel.

Rarezas del archivo real:
- "N/A" en toda pregunta que no aplicaba: se guarda como ausente.
- Las selecciones multiples vienen como texto JSON ('["Lunes"]').
- Fecha "2026/08/06 00:00:00" y horas "09:03:22", de Caracas.
- (0, 0) es "sin ubicacion"; si x/y quedan fuera de Venezuela pero al
  revez caen adentro, estaban invertidas.
- Las fotos son enlaces al servidor de Survey123: se guardan como enlace.
- El formulario de entonces no exigia empresa ni codigo a un cliente ya
  existente: esos censos entran, pero no se pueden cruzar con la maestra.

Reimportar es seguro: cada censo se identifica por su smid (Censo.externoId)
y los que ya estan no se tocan.
"""

from __future__ import annotations

import csv
import io
import json
import re
import secrets
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
from fastapi import HTTPException
from psycopg.types.json import Jsonb

from app.core.auth import hash_password
from app.core.ubicacion import fuera_de_venezuela, sin_ubicacion
from app.services import censo_clientes as cc
from app.services import formulario as fm
from app.services import xlsform as xf
from app.services.maestra_import import nombre_comparable

PREFIJO_EXTERNO = "survey123:"
VERSION = "survey123"
# Los censos de un encuestador que no es usuario de la app quedan a nombre de
# este usuario tecnico (inactivo: no puede entrar).
USUARIO_TECNICO = {"id": "usr-survey123", "nombre": "Importado de Survey123", "email": "survey123@importado.invalid"}
_DESFASE_CARACAS = timedelta(hours=-4)
_PUNTO_WKT = re.compile(r"POINT\s*\(\s*([-\d.eE]+)\s+([-\d.eE]+)\s*\)", re.IGNORECASE)
_DURACION = re.compile(r"^(\d{1,2}):(\d{2}):(\d{2})$")
_HORA = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")
MAXIMO_EJEMPLOS = 50


@dataclass
class Fila:
    externo_id: str
    respuestas: dict
    lat: float | None
    lng: float | None
    iniciado: datetime
    terminado: datetime
    invertida: bool = False
    # Traia coordenadas, pero no sirven (fuera de Venezuela, o x = y).
    invalida: bool = False


@dataclass
class Lectura:
    filas: list[Fila]
    errores: list[str] = field(default_factory=list)


# ---------- Lectura del archivo ----------


def _decodificar(contenido: bytes) -> str:
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    raise HTTPException(400, "No pude leer el texto del archivo (codificación desconocida)")


def _filas_crudas(contenido: bytes, nombre: str) -> list[dict]:
    """Del archivo a una lista de {pregunta: valor}, con los nombres en
    minuscula (el GeoJSON los trae en mayuscula)."""
    nombre = nombre.lower()
    if nombre.endswith(".xlsx"):
        try:
            libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, "No se pudo leer el Excel") from None
        filas = list(libro[libro.sheetnames[0]].iter_rows(values_only=True))
        libro.close()
        if not filas:
            return []
        encabezado = [str(v or "").strip().lower() for v in filas[0]]
        return [dict(zip(encabezado, f)) for f in filas[1:] if any(v is not None for v in f)]
    texto = _decodificar(contenido)
    if nombre.endswith((".json", ".geojson")) or texto.lstrip().startswith("{"):
        try:
            datos = json.loads(texto)
        except json.JSONDecodeError:
            raise HTTPException(400, "El JSON no es válido") from None
        if not isinstance(datos, dict) or not isinstance(datos.get("features"), list):
            raise HTTPException(400, "El JSON no es una exportación de Survey123 (falta 'features')")
        return [
            {str(k).lower(): v for k, v in (f.get("properties") or {}).items()}
            for f in datos["features"]
            if isinstance(f, dict)
        ]
    primera = texto.split("\n", 1)[0]
    separador = ";" if primera.count(";") > primera.count(",") else ","
    lector = csv.DictReader(io.StringIO(texto), delimiter=separador)
    return [{(k or "").strip().lower(): v for k, v in fila.items()} for fila in lector]


def _vacio(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and valor.strip().upper() in cc.SIN_DATO)


def _numero(valor: object) -> float | None:
    if _vacio(valor):
        return None
    try:
        return float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None


def _ubicacion(cruda: dict) -> tuple[float | None, float | None, bool, bool]:
    """(lat, lng, invertida, invalida). x es la longitud e y la latitud."""
    x, y = _numero(cruda.get("x")), _numero(cruda.get("y"))
    if x is None or y is None or fuera_de_venezuela(y, x) and fuera_de_venezuela(x, y):
        # x/y faltan o vienen rotas (exportaciones viejas perdian el punto
        # decimal): la columna WKT trae lo mismo, bien.
        m = _PUNTO_WKT.search(str(cruda.get("ubicacion") or ""))
        if m:
            x, y = float(m.group(1)), float(m.group(2))
    if x is None or y is None or sin_ubicacion(y, x):
        return None, None, False, False
    if not fuera_de_venezuela(y, x):
        return y, x, False, False
    if not fuera_de_venezuela(x, y):
        return x, y, True, False
    # Por ejemplo x = y = 10,46: el telefono perdio la longitud.
    return None, None, False, True


def _fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()[:10].replace("/", "-")
    try:
        return date.fromisoformat(texto)
    except ValueError:
        return None


def _hora(valor: object) -> tuple[int, int, int] | None:
    if hasattr(valor, "hour"):
        return valor.hour, valor.minute, valor.second
    m = _HORA.match(str(valor or "").strip())
    if not m or int(m.group(1)) > 23:
        return None
    return int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)


def _instantes(cruda: dict) -> tuple[datetime, datetime] | None:
    """Inicio y fin en UTC sin zona. La exportacion trae fecha y hora de
    inicio (de Caracas) y la hora de fin o la duracion."""
    dia, hora = _fecha(cruda.get("fecha")), _hora(cruda.get("hora"))
    if not dia or not hora:
        return None
    inicio = datetime(dia.year, dia.month, dia.day, *hora) - _DESFASE_CARACAS
    fin = None
    hora_fin = _hora(cruda.get("horafinent"))
    if hora_fin:
        candidato = datetime(dia.year, dia.month, dia.day, *hora_fin) - _DESFASE_CARACAS
        if inicio <= candidato <= inicio + timedelta(hours=12):
            fin = candidato
    if fin is None:
        m = _DURACION.match(str(cruda.get("horaduracionencuesta") or "").strip())
        fin = inicio + (timedelta(hours=int(m[1]), minutes=int(m[2]), seconds=int(m[3])) if m else timedelta())
    return inicio, fin


def _valor(pregunta: dict, valor: object) -> object:
    """Una respuesta con el tipo de su pregunta, o None si no aplicaba."""
    if _vacio(valor):
        return None
    tipo = pregunta["type"]
    if tipo == "select_multiple":
        if isinstance(valor, list):
            opciones = valor
        else:
            texto = str(valor).strip()
            try:
                opciones = json.loads(texto) if texto.startswith("[") else re.split(r"\s*,\s*", texto)
            except json.JSONDecodeError:
                opciones = re.split(r"\s*,\s*", texto.strip("[]"))
        return [str(o).strip() for o in opciones if str(o).strip()] or None
    if tipo == "integer":
        numero = _numero(valor)
        return int(numero) if numero is not None else None
    if tipo == "decimal":
        return _numero(valor)
    if tipo == "date":
        dia = _fecha(valor)
        return dia.isoformat() if dia else None
    if tipo == "time":
        hora = _hora(valor)
        return f"{hora[0]:02d}:{hora[1]:02d}:{hora[2]:02d}" if hora else None
    if tipo == "image":
        # Enlace a la foto en el servidor de Survey123 (solo https).
        texto = str(valor).strip()
        return texto if texto.lower().startswith("https://") else None
    texto = " ".join(str(valor).split())
    return texto or None


def leer(contenido: bytes, nombre: str) -> Lectura:
    crudas = _filas_crudas(contenido, nombre)
    if not crudas:
        raise HTTPException(400, "El archivo no trae filas")
    columnas = set().union(*(c.keys() for c in crudas))
    faltan = [c for c in ("tipo_cliente", "nombre_comercio", "fecha", "hora") if c not in columnas]
    if "smid" not in columnas and "id" not in columnas:
        faltan.append("smid")
    if faltan:
        raise HTTPException(
            400, f"No parece una exportación del censo de Survey123: faltan las columnas {', '.join(faltan)}"
        )

    preguntas = [p for p in xf.preguntas_de(fm.definicion(cc.FORMULARIO)) if p["type"] not in ("note",)]
    filas, errores, vistos = [], [], set()
    for numero, cruda in enumerate(crudas, start=2):
        externo = str(cruda.get("smid") or cruda.get("id") or "").strip()
        if externo.endswith(".0"):
            externo = externo[:-2]
        if not externo:
            errores.append(f"Fila {numero}: sin smid")
            continue
        if externo in vistos:
            errores.append(f"Fila {numero}: el smid {externo} está repetido en el archivo")
            continue
        instantes = _instantes(cruda)
        if not instantes:
            errores.append(f"Fila {numero} (smid {externo}): sin fecha u hora válidas")
            continue
        vistos.add(externo)
        lat, lng, invertida, invalida = _ubicacion(cruda)
        respuestas = {}
        for p in preguntas:
            if p["type"] == "geopoint":
                if lat is not None:
                    respuestas[p["name"]] = {"lat": lat, "lng": lng}
                continue
            valor = _valor(p, cruda.get(p["name"]))
            if valor is not None:
                respuestas[p["name"]] = valor
        filas.append(
            Fila(PREFIJO_EXTERNO + externo, respuestas, lat, lng, *instantes, invertida=invertida, invalida=invalida)
        )
    return Lectura(filas=filas, errores=errores)


# ---------- Importacion ----------


def _usuario_tecnico(cur) -> str:
    cur.execute('SELECT "id" FROM "Usuario" WHERE "id" = %s', (USUARIO_TECNICO["id"],))
    if not cur.fetchone():
        cur.execute(
            'INSERT INTO "Usuario" ("id", "nombre", "email", "passwordHash", "rol", "activo") '
            "VALUES (%s, %s, %s, %s, 'VENDEDOR', false)",
            (
                USUARIO_TECNICO["id"], USUARIO_TECNICO["nombre"], USUARIO_TECNICO["email"],
                hash_password(secrets.token_urlsafe(32)),
            ),
        )
    return USUARIO_TECNICO["id"]


def _usuarios_por_nombre(cur) -> dict[str, str]:
    cur.execute(
        'SELECT "id", "nombre" FROM "Usuario" WHERE "rol" = \'VENDEDOR\' AND "id" <> %s', (USUARIO_TECNICO["id"],)
    )
    por_nombre: dict[str, list[str]] = {}
    for u in cur.fetchall():
        por_nombre.setdefault(nombre_comparable(u["nombre"]), []).append(u["id"])
    # Un nombre que comparten dos usuarios no se puede asignar.
    return {n: ids[0] for n, ids in por_nombre.items() if len(ids) == 1}


def importar(cur, lectura: Lectura) -> dict:
    """Guarda los censos nuevos y los procesa (duplicados y clientes), en
    orden cronologico. No hace commit: la vista previa usa la misma funcion
    y deshace todo."""
    cur.execute(
        'SELECT "externoId" FROM "Censo" WHERE "externoId" = ANY(%s)',
        ([f.externo_id for f in lectura.filas],),
    )
    ya = {r["externoId"] for r in cur.fetchall()}
    nuevas = sorted((f for f in lectura.filas if f.externo_id not in ya), key=lambda f: (f.terminado, f.externo_id))

    usuarios = _usuarios_por_nombre(cur)
    tecnico = _usuario_tecnico(cur) if nuevas else None
    sin_usuario: Counter[str] = Counter()
    ctx = cc.Contexto(cur)
    acciones: Counter[str] = Counter()
    resumen = Counter()
    revisar, duplicados = [], []
    for f in nuevas:
        r = f.respuestas
        encuestador = r.get("encuestador")
        usuario_id = usuarios.get(nombre_comparable(encuestador))
        if not usuario_id:
            sin_usuario[encuestador or "(sin encuestador)"] += 1
            usuario_id = tecnico
        censo_id = f"cen-{uuid.uuid4().hex[:16]}"
        columnas = {col: r.get(nombre) for col, nombre in cc.COLUMNAS_DE_RESPUESTA.items()}
        cur.execute(
            'INSERT INTO "Censo" ("id", "formulario", "version", "usuarioId", "respuestas", "lat", "lng", '
            '"encuestador", "nombreComercio", "tipoCliente", "empresa", "ruta", "iniciadoEn", "terminadoEn", '
            '"recibidoEn", "origen", "externoId") '
            "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'SURVEY123', %s)",
            (
                censo_id, cc.FORMULARIO, VERSION, usuario_id, Jsonb(r), f.lat, f.lng,
                columnas["encuestador"], columnas["nombreComercio"], columnas["tipoCliente"],
                columnas["empresa"], columnas["ruta"], f.iniciado, f.terminado, f.terminado, f.externo_id,
            ),
        )
        proceso = cc.procesar(cur, censo_id, ctx)
        smid = f.externo_id.removeprefix(PREFIJO_EXTERNO)
        resumen["sinUbicacion"] += f.lat is None and not f.invalida
        resumen["ubicacionInvalida"] += f.invalida
        resumen["invertidas"] += f.invertida
        if proceso["duplicadoDe"]:
            resumen["duplicados"] += 1
            if len(duplicados) < MAXIMO_EJEMPLOS:
                duplicados.append({"smid": smid, "nombre": r.get("nombre_comercio")})
        if proceso["posibleCliente"]:
            resumen["posiblesClientes"] += 1
        elif not proceso["clientes"]:
            resumen["sinCodigo"] += 1
        for c in proceso["clientes"]:
            acciones[c["accion"]] += 1
            if c["accion"] == "REVISAR" and len(revisar) < MAXIMO_EJEMPLOS:
                revisar.append({"smid": smid, "empresa": c["empresa"], "codigo": c["codigo"], "motivo": c["motivo"]})

    return {
        "filas": len(lectura.filas) + len(lectura.errores),
        "errores": lectura.errores[:MAXIMO_EJEMPLOS],
        "totalErrores": len(lectura.errores),
        "yaImportadas": len(ya),
        "nuevas": len(nuevas),
        "sinUbicacion": resumen["sinUbicacion"],
        "ubicacionInvalida": resumen["ubicacionInvalida"],
        "invertidas": resumen["invertidas"],
        "duplicados": resumen["duplicados"],
        "ejemplosDuplicados": duplicados,
        "posiblesClientes": resumen["posiblesClientes"],
        "sinCodigo": resumen["sinCodigo"],
        "clientes": {a: acciones[a] for a in ("CREADO", "COMPLETADO", "EXISTENTE", "REVISAR")},
        "revisar": revisar,
        "encuestadoresSinUsuario": dict(sin_usuario.most_common()),
    }
