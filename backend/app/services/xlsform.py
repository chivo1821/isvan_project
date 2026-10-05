"""Lector de XLSForm (la plantilla de ArcGIS Survey123 Connect).

Convierte las hojas survey, choices y settings en la definicion que usa la
app (backend/app/formularios/<id>.json) para mostrar el formulario y para
validar cada envio. Se corre con backend/scripts/importar_formulario.py:
cuando cambie la encuesta, se vuelve a importar el Excel y no hay que tocar
codigo.

Respeta cada parametro de la fila: los que el motor usa quedan con nombre
propio (type, name, label, hint, guidance_hint, appearance, required,
required_message, readonly, default, calculation, constraint,
constraint_message, relevant, choice_filter y el largo de bind::esri:
fieldLength) y, ademas, TODAS las columnas no vacias de la fila se guardan
tal cual en `xlsform`, para que ninguno se pierda aunque el motor no lo use.

Lo que el motor no soporta (repeats, select_one_from_file, funciones que el
evaluador no conoce...) hace fallar la importacion con un mensaje claro, en
vez de quedar mal mostrado en el telefono del vendedor.
"""

from __future__ import annotations

import hashlib
import json
import re
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.services import xlsform_expresiones as ex

# Tipos de pregunta que el motor sabe mostrar y validar.
TIPOS = {
    "text", "integer", "decimal", "date", "time", "dateTime", "geopoint", "image",
    "note", "calculate", "hidden", "select_one", "select_multiple",
}
SI = {"yes", "true", "true()", "si", "sí", "1"}

# Columnas de la hoja survey que el motor usa, con el nombre que llevan en la
# definicion. Las demas solo quedan en `xlsform`.
COLUMNAS = {
    "name": "name",
    "label": "label",
    "hint": "hint",
    "guidance_hint": "guidance_hint",
    "appearance": "appearance",
    "required_message": "required_message",
    "default": "default",
    "calculation": "calculation",
    "constraint": "constraint",
    "constraint_message": "constraint_message",
    "relevant": "relevant",
    "choice_filter": "choice_filter",
}
# El largo maximo de un texto, con los dos nombres que usa Survey123.
COLUMNAS_LARGO = ("bind::esri:fieldLength", "bind_fieldLength")

# Defaults que no son un valor fijo sino "el dia / la hora al abrir".
DEFAULTS_DINAMICOS = {"today()", "now()"}


class ErrorXlsform(ValueError):
    pass


# ---------- HTML de las ayudas ----------
#
# Survey123 permite un poco de HTML en label, hint y guidance_hint (la
# encuesta usa <font color>, <b>, <i> y <h3 style="color: red">). Se deja
# pasar solo eso: etiquetas de formato y el color. Nada de scripts, enlaces
# ni atributos de eventos.

_ETIQUETAS = {"b", "strong", "i", "em", "u", "br", "p", "span", "font", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li"}
_COLOR = re.compile(r"^\s*(#[0-9a-fA-F]{3,8}|[a-zA-Z]{3,20}|rgb\(\s*\d+\s*,\s*\d+\s*,\s*\d+\s*\))\s*$")


class _Limpiador(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.salida: list[str] = []
        self.abiertas: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag not in _ETIQUETAS:
            return
        color = None
        for nombre, valor in attrs:
            if nombre == "color" and valor and _COLOR.match(valor):
                color = valor.strip()
            elif nombre == "style" and valor:
                m = re.search(r"(?:^|;)\s*color\s*:\s*([^;]+)", valor)
                if m and _COLOR.match(m.group(1)):
                    color = m.group(1).strip()
        if tag == "br":
            self.salida.append("<br>")
            return
        etiqueta = "span" if tag == "font" else tag
        self.salida.append(f'<{etiqueta} style="color: {color}">' if color else f"<{etiqueta}>")
        self.abiertas.append(etiqueta)

    def handle_endtag(self, tag):
        etiqueta = "span" if tag == "font" else tag
        if etiqueta in self.abiertas:
            # Cierra en orden aunque el HTML original no lo haga.
            while self.abiertas:
                ultima = self.abiertas.pop()
                self.salida.append(f"</{ultima}>")
                if ultima == etiqueta:
                    break

    def handle_data(self, data):
        self.salida.append(data.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    def resultado(self) -> str:
        while self.abiertas:
            self.salida.append(f"</{self.abiertas.pop()}>")
        return "".join(self.salida).strip()


def limpiar_html(texto: str | None) -> str | None:
    if texto is None:
        return None
    limpiador = _Limpiador()
    limpiador.feed(str(texto))
    limpiador.close()
    return limpiador.resultado() or None


def tiene_html(texto: str | None) -> bool:
    return bool(texto) and re.search(r"<[a-zA-Z/]", str(texto)) is not None


# ---------- Lectura ----------


def _filas(hoja) -> list[dict]:
    encabezados = [str(c.value).strip() if c.value is not None else "" for c in hoja[1]]
    filas = []
    for numero, valores in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        fila = {}
        for encabezado, valor in zip(encabezados, valores):
            if not encabezado or valor is None:
                continue
            texto = str(valor).strip() if not isinstance(valor, (int, float)) else valor
            if texto == "":
                continue
            fila[encabezado] = texto
        if fila:
            fila["_fila"] = numero
            filas.append(fila)
    return filas


def _si(valor) -> bool:
    return str(valor).strip().lower() in SI


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor)
    if isinstance(valor, float) and valor.is_integer():
        texto = str(int(valor))
    return texto


def _validar_expresion(texto: str, donde: str, nombres_validos: set[str], sueltos_validos: set[str]) -> None:
    try:
        nodo = ex.analizar(texto)
    except ex.ErrorExpresion as e:
        raise ErrorXlsform(f"{donde}: {e} (expresion: {texto!r})") from e
    faltan = ex.referencias(nodo) - nombres_validos
    if faltan:
        raise ErrorXlsform(f"{donde}: nombra preguntas que no existen: {sorted(faltan)}")
    sueltos = ex.nombres(nodo) - sueltos_validos
    if sueltos:
        raise ErrorXlsform(f"{donde}: nombres no reconocidos en la expresion: {sorted(sueltos)}")


def leer_xlsform(origen: str | Path | bytes, form_id: str | None = None) -> dict:
    """La definicion del formulario a partir del Excel."""
    libro = load_workbook(BytesIO(origen) if isinstance(origen, bytes) else origen, data_only=True)
    for hoja in ("survey", "choices"):
        if hoja not in libro.sheetnames:
            raise ErrorXlsform(f"Falta la hoja '{hoja}'")

    # Opciones
    listas: dict[str, list[dict]] = {}
    for fila in _filas(libro["choices"]):
        lista = _texto(fila.get("list_name"))
        nombre = _texto(fila.get("name"))
        if not lista or nombre is None:
            continue
        extra = {
            k: _texto(v)
            for k, v in fila.items()
            if k not in ("list_name", "name", "label", "_fila") and not k.startswith("media")
        }
        listas.setdefault(lista, []).append(
            {"name": nombre, "label": limpiar_html(_texto(fila.get("label"))) or nombre, "extra": extra}
        )
    # El mismo name puede repetirse si cambia alguna columna de filtro (la
    # parroquia "Sucre" del municipio Libertador y la del municipio Sucre):
    # el choice_filter las separa, como en Survey123. Repetida en todo, es
    # un error del Excel.
    for lista, opciones in listas.items():
        vistos = set()
        for o in opciones:
            clave = (o["name"], tuple(sorted(o["extra"].items())))
            if clave in vistos:
                raise ErrorXlsform(f"La lista '{lista}' repite la opcion '{o['name']}' con las mismas columnas")
            vistos.add(clave)

    # Preguntas
    filas = _filas(libro["survey"])
    raiz: dict = {"type": "group", "name": None, "children": []}
    pila = [raiz]
    planas: list[dict] = []
    for fila in filas:
        tipo_crudo = _texto(fila.get("type")) or ""
        partes = tipo_crudo.split()
        tipo = partes[0] if partes else ""
        donde = f"survey fila {fila['_fila']}"
        if tipo in ("begin_group", "begin group"):
            grupo = {
                "type": "group",
                "name": _texto(fila.get("name")),
                "label": limpiar_html(_texto(fila.get("label"))),
                "relevant": _texto(fila.get("relevant")),
                "appearance": _texto(fila.get("appearance")),
                "children": [],
                "xlsform": {k: _texto(v) for k, v in fila.items() if k != "_fila"},
            }
            pila[-1]["children"].append(grupo)
            pila.append(grupo)
            planas.append(grupo)
            continue
        if tipo in ("end_group", "end group"):
            if len(pila) == 1:
                raise ErrorXlsform(f"{donde}: end_group sin begin_group")
            pila.pop()
            continue
        if tipo.startswith("begin_repeat") or tipo.startswith("begin repeat"):
            raise ErrorXlsform(f"{donde}: los repeat todavia no estan soportados")
        if tipo not in TIPOS:
            raise ErrorXlsform(f"{donde}: tipo de pregunta no soportado: {tipo_crudo!r}")

        pregunta: dict = {"type": tipo}
        if tipo in ("select_one", "select_multiple"):
            if len(partes) != 2:
                raise ErrorXlsform(f"{donde}: '{tipo_crudo}' tiene que nombrar su lista")
            if partes[1] not in listas:
                raise ErrorXlsform(f"{donde}: la lista '{partes[1]}' no existe en choices")
            pregunta["list"] = partes[1]
        for columna, clave in COLUMNAS.items():
            valor = _texto(fila.get(columna))
            if valor is not None:
                pregunta[clave] = valor
        for columna in ("label", "hint", "guidance_hint", "required_message", "constraint_message"):
            if columna in pregunta:
                pregunta[columna] = limpiar_html(pregunta[columna])
        pregunta["required"] = _si(fila.get("required", ""))
        pregunta["readonly"] = _si(fila.get("readonly", ""))
        largo = next((fila[c] for c in COLUMNAS_LARGO if c in fila), None)
        if largo is not None:
            try:
                pregunta["field_length"] = int(float(largo))
            except ValueError as e:
                raise ErrorXlsform(f"{donde}: el largo '{largo}' no es un numero") from e
        if not pregunta.get("name"):
            raise ErrorXlsform(f"{donde}: la pregunta no tiene name")
        pregunta["xlsform"] = {k: _texto(v) for k, v in fila.items() if k != "_fila"}
        pila[-1]["children"].append(pregunta)
        planas.append(pregunta)

    if len(pila) != 1:
        raise ErrorXlsform(f"El grupo '{pila[-1]['name']}' no se cierra (falta end_group)")

    # Nombres unicos y expresiones validas
    preguntas = [p for p in planas if p["type"] != "group"]
    nombres = [p["name"] for p in preguntas]
    repetidos = {n for n in nombres if nombres.count(n) > 1}
    if repetidos:
        raise ErrorXlsform(f"Nombres de pregunta repetidos: {sorted(repetidos)}")
    validos = set(nombres)
    especiales = {"end_hour", "end_hh_mm_ss"}
    for p in planas:
        etiqueta = f"pregunta '{p['name']}'"
        if p.get("relevant"):
            _validar_expresion(p["relevant"], f"{etiqueta}, relevant", validos, set())
        if p.get("constraint"):
            _validar_expresion(p["constraint"], f"{etiqueta}, constraint", validos, {"string_length"})
        if p.get("calculation"):
            _validar_expresion(p["calculation"], f"{etiqueta}, calculation", validos, especiales)
        if p.get("default") and p["default"] in DEFAULTS_DINAMICOS:
            p["default_dinamico"] = p.pop("default")
        if p.get("choice_filter"):
            columnas = {"name", "label"} | {k for o in listas[p["list"]] for k in o["extra"]}
            _validar_expresion(p["choice_filter"], f"{etiqueta}, choice_filter", validos, columnas)

    settings = {}
    if "settings" in libro.sheetnames:
        filas_settings = _filas(libro["settings"])
        if filas_settings:
            settings = {k: _texto(v) for k, v in filas_settings[0].items() if k != "_fila"}

    definicion = {
        "id": form_id or settings.get("form_id") or "formulario",
        "title": settings.get("form_title") or form_id or "Formulario",
        "settings": settings,
        "children": raiz["children"],
        "choices": listas,
    }
    # La version es la huella de la definicion: cada censo guarda con cual se
    # lleno, y una encuesta reimportada sin cambios da la misma version.
    huella = hashlib.sha256(json.dumps(definicion, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    definicion["version"] = settings.get("version") or huella[:12]
    return definicion


def preguntas_de(definicion: dict) -> list[dict]:
    """Todas las preguntas en orden, con el grupo al que pertenecen."""
    salida: list[dict] = []

    def recorrer(nodos: list[dict], grupos: list[dict]):
        for nodo in nodos:
            if nodo["type"] == "group":
                recorrer(nodo["children"], grupos + [nodo])
            else:
                salida.append({**nodo, "_grupos": grupos})

    recorrer(definicion["children"], [])
    return salida
