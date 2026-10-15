"""Lectura de la maestra de clientes (el padron que el cliente mantiene
aparte del sistema de ventas), una por empresa.

Cada archivo es una FOTO completa: la carga nueva reemplaza a la anterior
(ver app/api/maestra.py). Sirve para saber que clientes faltan por censar y
para contrastar lo censado; la parte de activos (codigo, tipo y marca de la
nevera) se guarda pero todavia no se usa.

Sobre el archivo real (MAESTRA DE CLIENTES TRALOG, oct-2026): antes del
encabezado trae unas filas con el nombre de la empresa, su RIF y el total de
clientes, asi que el encabezado se busca; los encabezados vienen con puntos y
espacios de sobra ("Razón Social.", "TP ", "Frec."), hay filas repetidas
identicas y el encabezado se repite en medio de los datos.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter
from dataclasses import astuple, dataclass, field
from difflib import SequenceMatcher

import openpyxl
from fastapi import HTTPException

from app.core.excel_utils import mapear_columnas, normalizar_encabezado, valor_a_texto
from app.services.indicadores_venta import ruta_comparable

# campo -> encabezados que lo pueden nombrar (ya normalizados: minusculas,
# sin acentos y con "_" en vez de espacios, ver normalizar_encabezado).
ALIAS_COLUMNAS: dict[str, list[str]] = {
    "codigo": ["codigo", "cod", "codigo_cliente", "cod_cliente"],
    "razonSocial": ["razon_social.", "razon_social", "razon", "nombre", "cliente"],
    "ruta": ["ruta", "ruta_preventa"],
    "diaVisita": ["dia", "dia_visita", "dia_de_visita"],
    "frecuencia": ["frec.", "frec", "frecuencia"],
    "secuencia": ["sec.", "sec", "secuencia"],
    "direccion": ["direccion", "direccion_de_entrega", "direccion_entrega", "direccion_fiscal"],
    "condicion": ["condicion", "condicion_de_pago"],
    "tipoCliente": ["tp", "tipo", "tipo_de_cliente", "tipo_de_clientes", "tipo_cliente"],
    "listaPrecio": ["precio", "lista_precio", "lista_de_precio"],
    "zona": ["zona", "sector"],
    "observacion": ["observacion", "observaciones", "nota"],
    "activoCodigo": ["codigo_activo", "activo", "serial_activo"],
    "activoTipo": ["nevera_propia", "tipo_de_activo", "tipo_activo"],
    "activoMarca": ["marca", "modelo_de_activo", "modelo_activo"],
}
REQUERIDAS = ["codigo", "razonSocial"]
CAMPOS = [c for c in ALIAS_COLUMNAS if c != "codigo"]
CAMPOS_DE_ACTIVO = ("activoCodigo", "activoTipo", "activoMarca")

# El encabezado se busca en las primeras filas: el archivo real lo tiene en
# la 7 (antes, nombre de la empresa, RIF y total).
FILAS_ANTES_DEL_ENCABEZADO = 20
# Mas alla no hay datos; hay archivos con miles de columnas vacias.
COLUMNAS_UTILES = 40
# La ruta de los dados de baja: no se cuenta como "por censar".
RUTA_DE_BAJA = "77"

EMPRESAS = ("ISVAN", "TRALOG")


@dataclass
class FilaMaestra:
    codigo: str
    razonSocial: str
    ruta: str | None = None
    diaVisita: str | None = None
    frecuencia: str | None = None
    secuencia: str | None = None
    direccion: str | None = None
    condicion: str | None = None
    tipoCliente: str | None = None
    listaPrecio: str | None = None
    zona: str | None = None
    observacion: str | None = None
    activoCodigo: str | None = None
    activoTipo: str | None = None
    activoMarca: str | None = None


@dataclass
class Maestra:
    filas: list[FilaMaestra]
    hoja: str = ""
    fila_encabezado: int = 0
    repetidas_identicas: int = 0
    # Un cliente con varias neveras viene en varias filas, una por activo.
    con_varios_activos: int = 0
    # codigo -> cuantas filas con datos del cliente distintos (queda la ultima)
    repetidos_distintos: dict[str, int] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)


# ---------- Normalizacion compartida (la usa tambien censo_clientes) ----------


def codigo_normalizado(valor: object) -> str:
    """Sin ceros a la izquierda: ventas los manda asi ("000324") y la maestra
    como numero (324). Igual que _sin_ceros de app/api/despachos.py."""
    crudo = valor_a_texto(valor).strip().upper()
    return crudo.lstrip("0") or ("0" if crudo else "")


def nombre_comparable(texto: str | None) -> str:
    """Nombre para comparar: sin acentos, mayusculas, sin C.A./S.A./F.P./
    S.R.L. ni signos. El mismo criterio del analisis de la maestra contra
    ventas (oct-2026)."""
    s = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode().upper()
    s = re.sub(r"\b(C\.?\s?A|S\.?\s?A|F\.?\s?P|S\.?\s?R\.?\s?L)\b\.?", " ", s)
    return " ".join(re.sub(r"[^A-Z0-9 ]", " ", s).split())


def parecido(a: str | None, b: str | None) -> float:
    """0 a 1: que tanto se parecen dos nombres."""
    na, nb = nombre_comparable(a), nombre_comparable(b)
    if not na or not nb:
        return 0.0
    return SequenceMatcher(None, na, nb).ratio()


# Desde aca, con el mismo codigo, se considera el mismo cliente.
MISMO_CLIENTE = 0.80


def mismo_cliente(a: str | None, b: str | None) -> bool:
    """Dos nombres del mismo cliente: muy parecidos, o uno contiene al otro
    (la maestra agrega la sucursal: "FARMACIA MESVOL" y "FARMACIA MESVOL
    SILENCIO")."""
    na, nb = nombre_comparable(a), nombre_comparable(b)
    if not na or not nb:
        return False
    corto, largo = sorted((na, nb), key=len)
    if len(corto) >= 4 and (" " + corto + " ") in (" " + largo + " "):
        return True
    return SequenceMatcher(None, na, nb).ratio() >= MISMO_CLIENTE


# Palabras que no distinguen a un negocio de otro: el rubro, la forma
# juridica, sus abreviaturas en la maestra ("PAN Y PAST", "FRIG", "COML") y
# conectores. Lo que queda es el nombre propio ("OFANTINA", "LAS BRISAS").
_GENERICAS = frozenset("""
    A AL C CA CIA COMPANIA DE DEL E EL EN F FP LA LAS LO LOS P S SA SRL RL Y
    ABASTO ABASTOS ABAST AGROPEC AGROPECUARIA ALMACEN AUTOMERC AUTOMERCADO BAKERY BAR BODEGA BODEGON CAFE
    CAFETERIA CARNICERIA CHARCUT CHARCUTERIA CLUB COMERCIAL COMERCIALIZADORA COML CORP CORPORACION DELI
    DELICATESES DELICATESSEN DISTR DISTRIBUCIONES DISTRIBUIDORA EMPRENDIMIENTO EXPRES EXPRESS FARMACIA
    FARMACIAS FOOD FOODS FRIG FRIGORIFICO FRUTERIA GRUPO HELADERIA HELADOS INV INVERSIONES KIOSCO KIOSKO
    LEGUMBRES LICORERIA LICORES LUNCH LUNCHERIA MARKET MERCADO MINI MINIMARKET MULTISERVICIOS PAN PAND
    PANADERIA PANDERIA PAST PASTELERIA PANIFICADORA REST RESTAURANT RESTAURANTE SERVICIOS SUPER
    SUPERMERCADO SUPERMERCADOS TIENDA VIVERES
    CENTRO ESTE GRAN NORTE OESTE SUR JOSE MARIA
""".split())
_ANIO = re.compile(r"^(19|20)\d\d$")


def _distintivas(nombre: str | None) -> list[str]:
    return [p for p in nombre_comparable(nombre).split() if p not in _GENERICAS and not _ANIO.match(p)]


def _misma_palabra(a: str, b: str) -> bool:
    if a == b:
        return len(a) >= 2
    corta, larga = sorted((a, b), key=len)
    if len(corta) < 3 or corta.isdigit() or larga.isdigit():
        return False
    # Errores de tipeo ("EDELVYSE" y "EDELVAYS", "OTTU" y "OTU") y
    # abreviaturas ("HIGU" por "HIGUEROTE").
    if SequenceMatcher(None, a, b).ratio() >= MISMO_CLIENTE:
        return True
    return len(corta) >= 4 and larga.startswith(corta)


def mismo_negocio(a: str | None, b: str | None) -> bool:
    """El nombre que escribe el vendedor en el censo (el comercial, a mano)
    contra la razon social de la maestra o de ventas: comparten una palabra
    distintiva. "panadería la ofantina" y "PAN Y PAST OFANTINA C A" son el
    mismo; "Licoreria Los Compadres" y "AUTOMERC LOS DOS COMPAÑEROS", no."""
    if mismo_cliente(a, b):
        return True
    pa, pb = _distintivas(a), _distintivas(b)
    if any(_misma_palabra(x, y) for x in pa for y in pb):
        return True
    # Palabras pegadas o partidas: "red vital" y "REDVITAL".
    juntas_a, juntas_b = "".join(pa), "".join(pb)
    return any(len(p) >= 5 and p in juntas_b for p in pa) or any(len(p) >= 5 and p in juntas_a for p in pb)


# ---------- Lectura ----------


_ERROR_EXCEL = re.compile(r"^#(N/A|REF!|VALUE!|DIV/0!|NAME\?|NULL!|NUM!)$", re.IGNORECASE)


def _texto(valor: object) -> str | None:
    limpio = " ".join(valor_a_texto(valor).split())
    if not limpio or _ERROR_EXCEL.match(limpio):
        return None
    return limpio


def _texto_activo(valor: object) -> str | None:
    # "Sin activo" no viene vacio sino como 0.
    limpio = _texto(valor)
    return None if limpio in (None, "0") else limpio


def _sin_activo(fila: FilaMaestra) -> tuple:
    return tuple(v for k, v in vars(fila).items() if k not in CAMPOS_DE_ACTIVO)


def _buscar_encabezado(filas: list[tuple]) -> int:
    """El indice de la fila de encabezado: la primera que nombra el codigo
    y la razon social."""
    for i, fila in enumerate(filas[:FILAS_ANTES_DEL_ENCABEZADO]):
        nombres = {normalizar_encabezado(v) for v in fila if v is not None}
        if any(a in nombres for a in ALIAS_COLUMNAS["codigo"]) and any(
            a in nombres for a in ALIAS_COLUMNAS["razonSocial"]
        ):
            return i
    # Sin encabezado reconocible: el error de mapear_columnas lista lo que
    # trae la primera fila, que es lo que el usuario necesita ver.
    mapear_columnas(filas[0] if filas else (), ALIAS_COLUMNAS, REQUERIDAS)
    raise HTTPException(400, "No encontre la fila de encabezado (Código, Razón Social...)")


def _empresa_del_archivo(filas: list[tuple], nombre_archivo: str) -> set[str]:
    """Que empresa dice ser el archivo, por las filas de arriba (nombre de la
    empresa) o el nombre del archivo."""
    texto = " ".join(valor_a_texto(v) for fila in filas for v in fila if v is not None) + " " + nombre_archivo
    return {e for e in EMPRESAS if re.search(rf"\b{e}\b", texto.upper())}


def leer_maestra(contenido: bytes, nombre_archivo: str, empresa: str) -> Maestra:
    """Del Excel a filas tipadas. Solo falla (400) si el archivo no se puede
    usar o es de la otra empresa; lo demas sale como avisos."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo: tiene que ser un Excel (.xlsx)") from None

    hoja = libro[libro.sheetnames[0]]
    crudas = [f[:COLUMNAS_UTILES] for f in hoja.iter_rows(values_only=True)]
    libro.close()
    inicio = _buscar_encabezado(crudas)

    dice_ser = _empresa_del_archivo(crudas[:inicio], nombre_archivo)
    if dice_ser and empresa not in dice_ser:
        raise HTTPException(
            400, f"El archivo parece ser de {' y '.join(sorted(dice_ser))}, no de {empresa}. Revisa la empresa elegida."
        )

    mapa = mapear_columnas(crudas[inicio], ALIAS_COLUMNAS, REQUERIDAS)

    def celda(fila: tuple, nombre: str) -> object:
        indice = mapa.get(nombre)
        return fila[indice] if indice is not None and indice < len(fila) else None

    filas: list[FilaMaestra] = []
    posicion: dict[str, int] = {}
    identicas = 0
    varios_activos: set[str] = set()
    distintos: Counter[str] = Counter()
    sin_codigo = 0
    for cruda in crudas[inicio + 1:]:
        if all(v is None or valor_a_texto(v) == "" for v in cruda):
            continue
        codigo = codigo_normalizado(celda(cruda, "codigo"))
        if normalizar_encabezado(codigo) in ALIAS_COLUMNAS["codigo"]:
            continue  # el encabezado repetido en medio de los datos
        razon = _texto(celda(cruda, "razonSocial"))
        if not codigo:
            if razon:
                sin_codigo += 1
            continue
        fila = FilaMaestra(
            codigo=codigo,
            razonSocial=razon or f"(sin nombre) {codigo}",
            **{
                c: (_texto_activo if c in CAMPOS_DE_ACTIVO else _texto)(celda(cruda, c))
                for c in CAMPOS
                if c != "razonSocial"
            },
        )
        if codigo in posicion:
            anterior = filas[posicion[codigo]]
            if astuple(anterior) == astuple(fila):
                identicas += 1
            elif _sin_activo(anterior) == _sin_activo(fila):
                varios_activos.add(codigo)
            else:
                distintos[codigo] += 1
                # Se queda la ultima: es lo que hace un padron al actualizarse.
                filas[posicion[codigo]] = fila
        else:
            posicion[codigo] = len(filas)
            filas.append(fila)

    if not filas:
        raise HTTPException(400, "El archivo no trae ninguna fila con código de cliente")

    avisos: list[str] = []
    if identicas:
        avisos.append(f"{identicas} fila(s) repetidas idénticas: se cargó una sola de cada una")
    if distintos:
        muestra = ", ".join(sorted(distintos, key=lambda c: (len(c), c))[:8])
        avisos.append(
            f"{len(distintos)} código(s) repetidos con datos distintos; "
            f"se quedó la última fila de cada uno ({muestra})"
        )
    if sin_codigo:
        avisos.append(f"{sin_codigo} fila(s) con nombre pero sin código: no se cargaron")
    sin_ruta = sum(1 for f in filas if not f.ruta)
    if sin_ruta:
        avisos.append(f"{sin_ruta} cliente(s) sin ruta")

    return Maestra(
        filas=filas,
        hoja=hoja.title.strip(),
        fila_encabezado=inicio + 1,
        repetidas_identicas=identicas,
        con_varios_activos=len(varios_activos),
        repetidos_distintos=dict(distintos),
        avisos=avisos,
    )


def es_de_baja(ruta: str | None) -> bool:
    return ruta_comparable(ruta) == RUTA_DE_BAJA


def _orden_ruta(ruta: str) -> tuple:
    return (not ruta.isdigit(), int(ruta) if ruta.isdigit() else 0, ruta)


def resumen_de(maestra: Maestra) -> dict:
    """Como quedo el padron: cuantos por ruta y tipo, y los de baja."""
    def conteo(obtener) -> dict[str, int]:
        c = Counter((obtener(f) or "(sin dato)").strip().upper() for f in maestra.filas)
        return dict(sorted(c.items(), key=lambda kv: -kv[1]))

    por_ruta = Counter(ruta_comparable(f.ruta) or "(sin ruta)" for f in maestra.filas)
    de_baja = sum(1 for f in maestra.filas if es_de_baja(f.ruta))
    return {
        "hoja": maestra.hoja,
        "filaEncabezado": maestra.fila_encabezado,
        "clientes": len(maestra.filas),
        "deBaja": de_baja,
        "activos": len(maestra.filas) - de_baja,
        "repetidasIdenticas": maestra.repetidas_identicas,
        "conVariosActivos": maestra.con_varios_activos,
        "repetidosDistintos": len(maestra.repetidos_distintos),
        # Por numero de ruta, en orden (1, 2, ..., 12, 77).
        "porRuta": dict(sorted(por_ruta.items(), key=lambda kv: _orden_ruta(kv[0]))),
        "porTipo": conteo(lambda f: f.tipoCliente),
        "avisos": maestra.avisos,
    }


def contraste_con_ventas(cur, empresa: str, maestra: Maestra) -> dict:
    """La maestra contra los clientes que facturaron (VentaCliente): cuantos
    codigos coinciden, cuantos son el mismo cliente por nombre, y los que
    estan en un lado y no en el otro."""
    cur.execute('SELECT "codigo", "nombre", "ruta" FROM "VentaCliente" WHERE "empresa" = %s', (empresa,))
    en_ventas = {codigo_normalizado(f["codigo"]): f for f in cur.fetchall()}
    codigos = {f.codigo: f for f in maestra.filas}

    comunes = [c for c in en_ventas if c in codigos]
    otro_nombre = []
    ruta_distinta = 0
    for c in comunes:
        venta, fila = en_ventas[c], codigos[c]
        if not mismo_cliente(venta["nombre"], fila.razonSocial):
            otro_nombre.append({"codigo": c, "nombreVentas": venta["nombre"], "nombreMaestra": fila.razonSocial})
        if ruta_comparable(venta["ruta"]) != ruta_comparable(fila.ruta):
            ruta_distinta += 1
    faltan = [
        {"codigo": c, "nombre": f["nombre"], "ruta": f["ruta"]}
        for c, f in en_ventas.items()
        if c not in codigos
    ]

    def orden(x: dict) -> tuple:
        return (len(x["codigo"]), x["codigo"])

    return {
        "clientesConVentas": len(en_ventas),
        "codigosEnComun": len(comunes),
        "mismoCliente": len(comunes) - len(otro_nombre),
        "nombreDistinto": sorted(otro_nombre, key=orden)[:50],
        "nombreDistintoTotal": len(otro_nombre),
        "rutaDistintaTotal": ruta_distinta,
        "conVentasSinMaestra": sorted(faltan, key=orden)[:50],
        "conVentasSinMaestraTotal": len(faltan),
        "soloEnLaMaestra": len(codigos) - len(comunes),
    }
