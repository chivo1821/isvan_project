import io
import uuid

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile

from app.core.auth import get_current_user, requiere_rol
from app.core.db import get_connection
from app.core.permisos import es_repartidor, vehiculo_asignado
from app.core.excel_utils import mapear_columnas, valor_a_texto
from app.core.ubicacion import sin_ubicacion
from app.schemas import (
    Cliente,
    ClienteCreate,
    ImportarClientesConfirmarRequest,
    ImportarClientesPreviewResponse,
)

router = APIRouter(prefix="/clientes", tags=["clientes"])


@router.get("/importar/plantilla")
def descargar_plantilla_clientes():
    """Plantilla descargable con los encabezados exactos que reconoce el
    importador (ver ALIAS_COLUMNAS mas abajo) y una fila de ejemplo, para que
    no haya que recordar el formato de memoria."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Clientes"
    hoja.append([
        "codigo", "nombre", "tipo", "direccion", "ciudad", "lat", "lng", "telefono", "email", "ruta",
    ])
    hoja.append([
        "2118", "Distribuidora Don Pepe", "Distribuidor", "Calle Real, Sector Los Ruices", "Caracas",
        10.4956, -66.8836, "+58 212 555 0102", "pedidos@donpepe.com", "R-07",
    ])
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(len(str(c.value)) for c in columna if c.value is not None)
        hoja.column_dimensions[letra].width = max(10, ancho + 2)

    buffer = io.BytesIO()
    libro.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_clientes.xlsx"'},
    )


# Un REPARTIDOR no tiene acceso a la cartera de clientes: solo a los
# clientes que son parada de la ruta de su vehiculo, que son los que su
# pantalla necesita nombrar y ubicar (ver app/core/permisos.py).
_CLIENTES_DE_MI_RUTA = (
    'SELECT DISTINCT c.* FROM "Cliente" c '
    'JOIN "Despacho" d ON d."destinoClienteId" = c."id" '
    'JOIN "Ruta" r ON r."id" = d."rutaId" '
    'WHERE r."vehiculoId" = %s'
)


@router.get("", response_model=list[Cliente])
def listar_clientes(usuario: dict = Depends(get_current_user)):
    with get_connection() as conn, conn.cursor() as cur:
        if es_repartidor(usuario):
            vehiculo_id = vehiculo_asignado(usuario)
            if not vehiculo_id:
                return []
            cur.execute(f'{_CLIENTES_DE_MI_RUTA} ORDER BY c."nombre"', (vehiculo_id,))
        else:
            cur.execute('SELECT * FROM "Cliente" ORDER BY "nombre"')
        return cur.fetchall()


@router.get("/{cliente_id}", response_model=Cliente)
def obtener_cliente(cliente_id: str, usuario: dict = Depends(get_current_user)):
    with get_connection() as conn, conn.cursor() as cur:
        if es_repartidor(usuario):
            vehiculo_id = vehiculo_asignado(usuario)
            if not vehiculo_id:
                raise HTTPException(404, "Cliente no encontrado")
            cur.execute(f'{_CLIENTES_DE_MI_RUTA} AND c."id" = %s', (vehiculo_id, cliente_id))
        else:
            cur.execute('SELECT * FROM "Cliente" WHERE "id" = %s', (cliente_id,))
        row = cur.fetchone()
    if not row:
        raise HTTPException(404, "Cliente no encontrado")
    return row


def _crear_cliente_interno(cur, data: ClienteCreate) -> dict:
    cliente_id = f"cli-{uuid.uuid4().hex[:10]}"
    # Unas coordenadas en (0, 0) son un dato faltante disfrazado: se guardan
    # como NULL para que el cliente aparezca claramente "sin ubicacion" en
    # vez de como un punto en medio del Atlantico.
    lat, lng = (None, None) if sin_ubicacion(data.lat, data.lng) else (data.lat, data.lng)
    cur.execute(
        'INSERT INTO "Cliente" '
        '("id", "empresa", "codigo", "nombre", "tipo", "direccion", "ciudad", "lat", "lng", "telefono", '
        '"email", "rutaComercial") '
        "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING *",
        (
            cliente_id, data.empresa, data.codigo, data.nombre, data.tipo,
            data.direccion, data.ciudad, lat, lng, data.telefono, data.email,
            data.rutaComercial,
        ),
    )
    return cur.fetchone()


@router.post("", response_model=Cliente, status_code=201, dependencies=[Depends(requiere_rol("DESPACHOS"))])
def crear_cliente(data: ClienteCreate):
    with get_connection() as conn, conn.cursor() as cur:
        # (empresa, codigo) es la llave de negocio real — el mismo codigo
        # puede referirse a clientes distintos en ISVAN y en TRALOG.
        cur.execute(
            'SELECT 1 FROM "Cliente" WHERE "empresa" = %s AND "codigo" = %s',
            (data.empresa, data.codigo),
        )
        if cur.fetchone():
            raise HTTPException(400, f'Ya existe un cliente con el codigo "{data.codigo}" en {data.empresa}')

        row = _crear_cliente_interno(cur, data)
        conn.commit()
        return row


# ---------- Carga masiva por Excel ----------
#
# Reglas confirmadas con el cliente: el codigo lo asigna el negocio, nunca se
# genera aca (por eso es obligatorio en el Excel, igual que nombre/tipo/
# direccion/ciudad); lat/lng son obligatorias (un cliente sin coordenadas no
# se puede incluir en una ruta); telefono es obligatorio (lo usan los
# despachadores para contactar al cliente); la empresa se elige una vez antes
# de subir el archivo (igual que en la importacion de despachos), no es una
# columna. Un codigo que ya existe en esa empresa se rechaza como error —
# nunca se sobreescribe un cliente existente por una carga masiva.
#
# La columna "ruta" (ruta comercial de venta/reparto a la que el negocio
# asigna al cliente) es opcional aca: el extracto de ventas tambien la trae y
# es la fuente de verdad — al importar despachos se refresca el valor del
# cliente (ver app/api/despachos.py). Se pide igual en esta plantilla para
# poder cargar clientes nuevos ya con su ruta, sin esperar a la primera
# venta.

ALIAS_COLUMNAS: dict[str, list[str]] = {
    "codigo": ["codigo", "codigo_cliente", "cod_cliente"],
    "nombre": ["nombre", "cliente", "razon_social"],
    "tipo": ["tipo", "tipo_cliente"],
    "direccion": ["direccion", "direccion_completa"],
    "ciudad": ["ciudad"],
    "lat": ["lat", "latitud"],
    "lng": ["lng", "lon", "longitud"],
    "telefono": ["telefono", "telefonos", "tel"],
    "email": ["email", "correo"],
    "ruta": ["ruta", "ruta_comercial", "cod_ruta", "codigo_ruta", "nro_ruta", "zona"],
}
CAMPOS_REQUERIDOS = ["codigo", "nombre", "tipo", "direccion", "ciudad", "lat", "lng", "telefono"]


@router.post(
    "/importar/preview",
    response_model=ImportarClientesPreviewResponse,
    dependencies=[Depends(requiere_rol("DESPACHOS"))],
)
def importar_clientes_preview(empresa: str = Form(...), archivo: UploadFile = File(...)):
    if empresa not in ("ISVAN", "TRALOG"):
        raise HTTPException(400, "empresa debe ser ISVAN o TRALOG")

    try:
        libro = openpyxl.load_workbook(io.BytesIO(archivo.file.read()), read_only=True, data_only=True)
        filas = list(libro.active.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo — verifica que sea un Excel (.xlsx) valido")

    if not filas:
        raise HTTPException(400, "El archivo esta vacio")

    mapa = mapear_columnas(filas[0], ALIAS_COLUMNAS, CAMPOS_REQUERIDOS)
    errores: list[dict] = []
    clientes_validos: list[dict] = []
    codigos_en_archivo: dict[str, int] = {}

    with get_connection() as conn, conn.cursor() as cur:
        # Los codigos ya existentes se traen de una sola vez, no una consulta
        # por fila del Excel (un archivo real trae miles de filas -> miles de
        # round-trips a la base, insoportable con Postgres en otra region).
        cur.execute('SELECT "codigo" FROM "Cliente" WHERE "empresa" = %s', (empresa,))
        codigos_existentes = {r["codigo"] for r in cur.fetchall()}

        for n, fila in enumerate(filas[1:], start=2):
            if fila is None or all(v is None for v in fila):
                continue  # fila vacia, se ignora

            def val(campo: str):
                idx = mapa.get(campo)
                return fila[idx] if idx is not None and idx < len(fila) else None

            codigo = valor_a_texto(val("codigo"))
            nombre = str(val("nombre") or "").strip()
            tipo = str(val("tipo") or "").strip()
            direccion = str(val("direccion") or "").strip()
            ciudad = str(val("ciudad") or "").strip()
            telefono = valor_a_texto(val("telefono"))
            email_raw = val("email")
            email = str(email_raw).strip() or None if email_raw is not None else None
            ruta_comercial = valor_a_texto(val("ruta")) or None

            error: tuple[str, str] | None = None
            if not codigo:
                error = ("codigo", "Falta el codigo de cliente")
            elif not nombre:
                error = ("nombre", "Falta el nombre")
            elif not tipo:
                error = ("tipo", "Falta el tipo de cliente")
            elif not direccion:
                error = ("direccion", "Falta la direccion")
            elif not ciudad:
                error = ("ciudad", "Falta la ciudad")
            elif not telefono:
                error = ("telefono", "Falta el telefono")

            lat = lng = None
            if error is None:
                try:
                    lat = float(val("lat"))
                    lng = float(val("lng"))
                    if not (-90 <= lat <= 90) or not (-180 <= lng <= 180):
                        raise ValueError
                except (TypeError, ValueError):
                    error = ("lat/lng", "lat y lng deben ser numeros validos (lat entre -90 y 90, lng entre -180 y 180)")
                else:
                    # (0, 0) cae en el golfo de Guinea: es un dato faltante
                    # cargado como cero, no una ubicacion real. Ver
                    # app/core/ubicacion.py.
                    if sin_ubicacion(lat, lng):
                        error = ("lat/lng", "Las coordenadas estan en 0,0 — falta la ubicacion real del cliente")

            if error is None and codigo in codigos_en_archivo:
                error = ("codigo", f'El codigo "{codigo}" esta repetido en la fila {codigos_en_archivo[codigo]}')

            if error is None and codigo in codigos_existentes:
                error = ("codigo", f'El codigo "{codigo}" ya existe en {empresa}')

            if error is not None:
                columna, motivo = error
                errores.append({"fila": n, "columna": columna, "motivo": motivo})
                continue

            codigos_en_archivo[codigo] = n
            clientes_validos.append({
                "empresa": empresa, "codigo": codigo, "nombre": nombre, "tipo": tipo,
                "direccion": direccion, "ciudad": ciudad, "lat": lat, "lng": lng,
                "telefono": telefono, "email": email, "rutaComercial": ruta_comercial,
            })

    errores.sort(key=lambda e: e["fila"])
    return {"clientes": clientes_validos, "errores": errores}


@router.post(
    "/importar/confirmar",
    response_model=list[Cliente],
    dependencies=[Depends(requiere_rol("DESPACHOS"))],
)
def importar_clientes_confirmar(data: ImportarClientesConfirmarRequest):
    if not data.clientes:
        raise HTTPException(400, "No hay clientes validos para importar")

    with get_connection() as conn, conn.cursor() as cur:
        # Revalida unicidad por si cambio algo entre el preview y la
        # confirmacion (ej. otra persona creo el mismo codigo mientras tanto)
        # -- en una consulta por empresa, no una por cliente.
        for empresa in {c.empresa for c in data.clientes}:
            codigos = [c.codigo for c in data.clientes if c.empresa == empresa]
            cur.execute(
                'SELECT "codigo" FROM "Cliente" WHERE "empresa" = %s AND "codigo" = ANY(%s)',
                (empresa, codigos),
            )
            ya_existentes = {r["codigo"] for r in cur.fetchall()}
            if ya_existentes:
                raise HTTPException(
                    400,
                    f'Estos codigos ya existen en {empresa} (creados por otra carga): {", ".join(sorted(ya_existentes))}',
                )

        creados = [_crear_cliente_interno(cur, c) for c in data.clientes]
        conn.commit()
        return creados
