"""Portal de autogestion de los clientes: su cartera, su estado de cuenta y
el reporte de sus pagos.

Cada endpoint toma los clientes de la cuenta de la sesion
(cuenta["clientes"], ver app/core/portal_auth.py) y no recibe del navegador
ningun dato que amplie eso: un cliente solo puede elegir entre los suyos.
"""

from __future__ import annotations

import re
import uuid
from datetime import date

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Path, Query, Request
from fastapi.responses import RedirectResponse, Response
from pydantic import BaseModel, Field

from app.api.reportes import _como_adjunto, _escribir_hoja
from app.core.db import get_connection
from app.core.fechas import hoy_caracas
from app.core.portal_auth import cuenta_de_sesion, get_cuenta_portal
from app.schemas import PagoPortalCreate, en_utc
from app.services import almacen_fotos as af
from app.services import bancos, cartera, documento_cobro, pagos, tasa_bcv

router = APIRouter(prefix="/portal", tags=["portal"])

COMPROBANTE_MAXIMO_BYTES = 10_000_000
TIPOS_COMPROBANTE = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp", "application/pdf": "pdf"}
_EXTENSION_TIPO = {v: k for k, v in TIPOS_COMPROBANTE.items()}
RUTA_CARGA_LOCAL = "/portal/pagos/comprobante/carga"


def _almacen():
    try:
        return af.almacen()
    except af.AlmacenNoConfigurado as e:
        raise HTTPException(503, str(e)) from e


def _prefijo(cuenta: dict) -> str:
    return "pagos/" + re.sub(r"[^A-Za-z0-9_-]", "_", cuenta["id"]) + "/"


def _cliente_elegido(cuenta: dict, cliente: str | None) -> list[cartera.Clave]:
    """"ISVAN:123" -> [("ISVAN", "123")] si es de la cuenta; sin elegir,
    todos los suyos."""
    if not cliente:
        return list(cuenta["clientes"])
    empresa, _, codigo = cliente.partition(":")
    clave = (empresa, codigo)
    if clave not in cuenta["clientes"]:
        raise HTTPException(404, "Cliente no encontrado")
    return [clave]


@router.get("/me")
def yo(cuenta: dict = Depends(cuenta_de_sesion)):
    """La cuenta, este o no validada (la pantalla dice "en validacion")."""
    with get_connection() as conn, conn.cursor() as cur:
        datos = cartera.datos_clientes(cur, cuenta["clientes"]) if cuenta["clientes"] else {}
    return {
        "id": cuenta["id"], "email": cuenta["email"], "nombre": cuenta["nombre"], "rif": cuenta["rif"],
        "telefono": cuenta["telefono"], "estado": cuenta["estado"],
        "clientes": [
            {"empresa": e, "codigoCliente": c, "nombre": (datos.get((e, c)) or {}).get("nombre") or f"Cliente {c}"}
            for e, c in cuenta["clientes"]
        ],
    }


@router.get("/resumen")
def resumen(cuenta: dict = Depends(get_cuenta_portal)):
    with get_connection() as conn, conn.cursor() as cur:
        p = cartera.parametros(cur)
        estados = cartera.estados(cur, cuenta["clientes"], p=p)
        datos = cartera.datos_clientes(cur, cuenta["clientes"])
        cur.execute(
            "SELECT count(*) AS \"cantidad\", COALESCE(sum(\"montoUsd\"), 0) AS \"total\" FROM \"Pago\" "
            "WHERE \"estado\" = 'REPORTADO' AND (\"empresa\"::text, \"codigoCliente\") IN "
            "(SELECT * FROM unnest(%s::text[], %s::text[]))",
            ([c[0] for c in cuenta["clientes"]], [c[1] for c in cuenta["clientes"]]),
        )
        en_revision = cur.fetchone()
    clientes = []
    for clave in cuenta["clientes"]:
        est = estados.get(clave) or cartera.EstadoCliente(*clave)
        clientes.append({**est.resumen(p), "nombre": (datos.get(clave) or {}).get("nombre") or f"Cliente {clave[1]}"})
    return {
        "clientes": clientes,
        "saldoUsd": sum(c["saldoUsd"] for c in clientes),
        "vencidoUsd": sum(c["vencidoUsd"] for c in clientes),
        "pagosEnRevision": en_revision["cantidad"],
        "pagosEnRevisionUsd": float(en_revision["total"]),
        "diasTolerancia": p.tolerancia,
        "bloqueoActivo": p.bloqueo_activo,
    }


@router.get("/documentos")
def documentos(
    cuenta: dict = Depends(get_cuenta_portal),
    estado: str = Query("pendientes", pattern="^(pendientes|todos)$"),
    cliente: str | None = Query(None, max_length=40),
):
    claves = _cliente_elegido(cuenta, cliente)
    with get_connection() as conn, conn.cursor() as cur:
        p = cartera.parametros(cur)
        estados = cartera.estados(cur, claves, p=p)
    salida = [d.como_dict(p.hoy) for est in estados.values() for d in est.documentos]
    if estado == "pendientes":
        salida = [d for d in salida if d["saldoUsd"] > 0]
    salida.sort(key=lambda d: (d["vence"], d["numero"]), reverse=estado == "todos")
    return salida


@router.get("/documentos/{documento_id}")
def documento(documento_id: str = Path(max_length=64), cuenta: dict = Depends(get_cuenta_portal)):
    """Una factura o nota de entrega del cliente: sus productos y sus abonos.
    Del despacho solo si ya se entrego y cuando (sin chofer, vehiculo ni
    ruta: eso es de la operacion interna)."""
    with get_connection() as conn, conn.cursor() as cur:
        d = documento_cobro.detalle(cur, documento_id)
    # 404 tambien si es de otro cliente: no revelar que existe.
    if d is None or (d["empresa"], d["codigoCliente"]) not in cuenta["clientes"]:
        raise HTTPException(404, "Documento no encontrado")
    despacho = d.pop("despacho")
    d.pop("rutaComercial", None)
    d.pop("origen", None)
    d["entrega"] = (
        {"estado": despacho["estado"], "entregadoEn": en_utc(despacho["entregadoEn"]) if despacho["entregadoEn"] else None}
        if despacho else None
    )
    return d


def _rango(desde: date | None, hasta: date | None) -> tuple[date | None, date | None]:
    if desde and hasta and desde > hasta:
        raise HTTPException(422, "La fecha inicial es posterior a la final")
    return desde, hasta


@router.get("/estado-cuenta")
def estado_cuenta(
    cuenta: dict = Depends(get_cuenta_portal),
    desde: date | None = None,
    hasta: date | None = None,
    cliente: str | None = Query(None, max_length=40),
):
    desde, hasta = _rango(desde, hasta)
    with get_connection() as conn, conn.cursor() as cur:
        return cartera.movimientos(cur, _cliente_elegido(cuenta, cliente), desde, hasta)


@router.get("/estado-cuenta.xlsx")
def estado_cuenta_excel(
    cuenta: dict = Depends(get_cuenta_portal),
    desde: date | None = None,
    hasta: date | None = None,
    cliente: str | None = Query(None, max_length=40),
):
    desde, hasta = _rango(desde, hasta)
    claves = _cliente_elegido(cuenta, cliente)
    with get_connection() as conn, conn.cursor() as cur:
        libro = estado_cuenta_libro(cur, claves, desde, hasta, cuenta["nombre"])
    return _como_adjunto(libro, "estado_de_cuenta")


def estado_cuenta_libro(cur, claves: list[cartera.Clave], desde: date | None, hasta: date | None, titular: str):
    """El Excel del estado de cuenta (lo usan el portal y cobranzas)."""
    p = cartera.parametros(cur)
    datos = cartera.datos_clientes(cur, claves)
    mov = cartera.movimientos(cur, claves, desde, hasta, p=p)
    estados = cartera.estados(cur, claves, p=p)
    nombre = lambda e, c: (datos.get((e, c)) or {}).get("nombre") or c  # noqa: E731

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Estado de cuenta"
    filas = [[None, None, None, "Saldo inicial", None, None, None, mov["saldoInicialUsd"]]] if desde else []
    filas += [
        [m["fecha"], m["empresa"], nombre(m["empresa"], m["codigoCliente"]), m["tipo"], m["referencia"],
         m["cargoUsd"] or None, m["abonoUsd"] or None, m["saldoUsd"]]
        for m in mov["movimientos"]
    ]
    _escribir_hoja(hoja, ["Fecha", "Empresa", "Cliente", "Tipo", "Documento", "Cargo USD", "Abono USD", "Saldo USD"], filas)
    for fila in hoja.iter_rows(min_row=2, min_col=6, max_col=8):
        for celda in fila:
            celda.number_format = "#,##0.00"

    pendientes = libro.create_sheet("Documentos pendientes")
    docs = sorted((d for est in estados.values() for d in est.pendientes()), key=lambda d: (d.vence, d.numero))
    _escribir_hoja(
        pendientes,
        ["Empresa", "Cliente", "Tipo", "Número", "Fecha", "Vence", "Monto USD", "Abonado USD", "Saldo USD", "Días de atraso"],
        [[d.empresa, nombre(d.empresa, d.codigoCliente), d.tipoDoc, d.numero, d.fecha, d.vence, float(d.montoUsd),
          float(d.montoUsd - d.saldoUsd), float(d.saldoUsd), d.dias_atraso(p.hoy) or None] for d in docs],
    )
    for fila in pendientes.iter_rows(min_row=2, min_col=7, max_col=9):
        for celda in fila:
            celda.number_format = "#,##0.00"
    portada = libro.create_sheet("Resumen", 0)
    portada.append(["Estado de cuenta"])
    portada.append(["Titular", titular])
    portada.append(["Emitido", hoy_caracas()])
    portada.append(["Período", f"{desde or 'inicio'} a {hasta or hoy_caracas()}"])
    portada.append([])
    portada.append(["Saldo USD", sum(float(e.saldo_usd()) for e in estados.values())])
    portada.append(["Vencido USD", sum(float(e.vencido_usd(p.hoy)) for e in estados.values())])
    portada.column_dimensions["A"].width = 18
    portada.column_dimensions["B"].width = 40
    portada["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    return libro


# ---------- Pagos ----------


@router.get("/pagos")
def mis_pagos(cuenta: dict = Depends(get_cuenta_portal)):
    with get_connection() as conn, conn.cursor() as cur:
        lista = pagos.listar(
            cur,
            '(p."empresa"::text, p."codigoCliente") IN (SELECT * FROM unnest(%(e)s::text[], %(c)s::text[]))',
            {"e": [c[0] for c in cuenta["clientes"]], "c": [c[1] for c in cuenta["clientes"]]},
        )
    # Al cliente no se le muestra quien de cobranzas lo reviso.
    return [
        {k: (en_utc(v) if k in ("reportadoEn", "revisadoEn") and v else v) for k, v in pago.items()
         if k not in ("revisadoPor", "registradoPor", "cuentaPortalId", "reportadoPorEmail")}
        for pago in lista
    ]


class SolicitudComprobante(BaseModel):
    tipo: str = Field(max_length=100)
    bytes: int = Field(gt=0)


@router.post("/pagos/comprobante/subida")
def pedir_subida_comprobante(datos: SolicitudComprobante, cuenta: dict = Depends(get_cuenta_portal)):
    """Enlace firmado para subir UN comprobante (imagen o PDF)."""
    tipo = datos.tipo.split(";")[0].strip().lower()
    if tipo not in TIPOS_COMPROBANTE:
        raise HTTPException(422, "El comprobante tiene que ser una imagen (JPG, PNG, WEBP) o un PDF")
    if datos.bytes > COMPROBANTE_MAXIMO_BYTES:
        raise HTTPException(422, f"El comprobante pesa más de {COMPROBANTE_MAXIMO_BYTES // 1_000_000} MB")
    clave = f"{_prefijo(cuenta)}{uuid.uuid4().hex}.{TIPOS_COMPROBANTE[tipo]}"
    subida = _almacen().subida(clave, tipo, datos.bytes, ruta_carga=RUTA_CARGA_LOCAL)
    return {"clave": clave, "url": subida.url, "metodo": subida.metodo, "cabeceras": subida.cabeceras,
            "conSesion": subida.con_sesion}


@router.put(RUTA_CARGA_LOCAL.removeprefix("/portal") + "/{ficha}", status_code=204)
async def carga_local(ficha: str, request: Request, cuenta: dict = Depends(get_cuenta_portal)):
    """Solo en desarrollo (ALMACEN_FOTOS=local): hace de R2."""
    almacen = _almacen()
    if not isinstance(almacen, af.AlmacenLocal):
        raise HTTPException(404, "No disponible")
    try:
        carga = almacen.leer_ficha(ficha)
    except ValueError as e:
        raise HTTPException(403, str(e)) from e
    if not carga["c"].startswith(_prefijo(cuenta)):
        raise HTTPException(403, "Ese enlace no es tuyo")
    if (request.headers.get("content-type") or "").split(";")[0].strip().lower() != carga["t"]:
        raise HTTPException(403, "El tipo del archivo no coincide con el enlace")
    contenido = await request.body()
    if len(contenido) != carga["b"]:
        raise HTTPException(403, "El tamaño del archivo no coincide con el enlace")
    almacen.guardar(carga["c"], contenido)


def _comprobante_valido(cur, cuenta: dict, clave: str) -> tuple[str, str]:
    """(clave, tipo) si el archivo esta subido, es de esta cuenta, no se uso
    en otro pago y respeta tipo y tamano."""
    try:
        af.validar_clave(clave)
    except ValueError:
        raise HTTPException(422, "Comprobante inválido") from None
    if not clave.startswith(_prefijo(cuenta)):
        raise HTTPException(422, "Ese comprobante no lo subiste tú")
    cur.execute('SELECT 1 FROM "Pago" WHERE "comprobanteClave" = %s', (clave,))
    if cur.fetchone():
        raise HTTPException(422, "Ese comprobante ya está en otro pago")
    guardada = _almacen().guardada(clave)
    if guardada is None:
        raise HTTPException(422, "El comprobante no terminó de subir; vuelve a intentar")
    tipo = guardada.tipo.split(";")[0].strip().lower() or _EXTENSION_TIPO.get(clave.rsplit(".", 1)[-1], "")
    if tipo not in TIPOS_COMPROBANTE or guardada.bytes > COMPROBANTE_MAXIMO_BYTES:
        raise HTTPException(422, "El comprobante no es válido (tipo o tamaño)")
    return clave, tipo


@router.post("/pagos", status_code=201)
def reportar_pago(data: PagoPortalCreate, cuenta: dict = Depends(get_cuenta_portal)):
    clave = (data.empresa, data.codigoCliente)
    if clave not in cuenta["clientes"]:
        raise HTTPException(404, "Cliente no encontrado")
    try:
        datos = pagos.normalizar(data.model_dump(), desde_portal=True)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    with get_connection() as conn, conn.cursor() as cur:
        comprobante = _comprobante_valido(cur, cuenta, data.comprobanteClave) if data.comprobanteClave else None
        try:
            pago_id = pagos.registrar(
                cur, empresa=data.empresa, codigo=data.codigoCliente, datos=datos,
                aplicaciones=[a.model_dump() for a in data.aplicaciones], cuenta_portal_id=cuenta["id"],
                comprobante=comprobante,
            )
        except ValueError as e:
            raise HTTPException(422, str(e)) from e
        conn.commit()
        return pagos.listar(cur, 'p."id" = %s', (pago_id,))[0]


@router.get("/pagos/{pago_id}/comprobante")
def ver_comprobante(pago_id: str, cuenta: dict = Depends(get_cuenta_portal)):
    with get_connection() as conn, conn.cursor() as cur:
        cur.execute(
            'SELECT "empresa"::text AS "empresa", "codigoCliente", "comprobanteClave", "comprobanteTipo" '
            'FROM "Pago" WHERE "id" = %s',
            (pago_id,),
        )
        fila = cur.fetchone()
    if not fila or (fila["empresa"], fila["codigoCliente"]) not in cuenta["clientes"] or not fila["comprobanteClave"]:
        raise HTTPException(404, "Comprobante no encontrado")
    return servir_comprobante(fila["comprobanteClave"], fila["comprobanteTipo"])


def servir_comprobante(clave: str, tipo: str | None):
    """R2: redirige a un enlace firmado de pocos minutos. Local: lo sirve."""
    almacen = _almacen()
    url = almacen.url_lectura(clave)
    if url:
        return RedirectResponse(url, status_code=302, headers={"Cache-Control": "private, no-store"})
    try:
        contenido = almacen.leer(clave)
    except FileNotFoundError:
        raise HTTPException(404, "El comprobante no está en el almacenamiento") from None
    return Response(contenido, media_type=tipo or "application/octet-stream", headers={"Cache-Control": "private, no-store"})


@router.get("/tasa")
def tasa(cuenta: dict = Depends(get_cuenta_portal)):
    """La ultima tasa BCV publicada, para mostrar el equivalente en USD
    mientras el cliente llena el reporte (el servidor la vuelve a tomar al
    guardar)."""
    with get_connection() as conn, conn.cursor() as cur:
        try:
            fila = tasa_bcv.tasa_actual(cur)
        except tasa_bcv.TasaNoDisponible as e:
            raise HTTPException(503, str(e)) from e
    return {"fecha": fila["fecha"], "tasa": float(fila["tasa"])}


@router.get("/cuentas-empresa")
def cuentas_empresa(cuenta: dict = Depends(get_cuenta_portal)):
    """Donde pagar: las cuentas activas de las empresas de sus clientes."""
    empresas = sorted({e for e, _ in cuenta["clientes"]})
    with get_connection() as conn, conn.cursor() as cur:
        cur.execute(
            'SELECT "id", "empresa"::text AS "empresa", "bancoCodigo", "tipo", "numero", "rif", "titular" '
            'FROM "CuentaBancariaEmpresa" WHERE "activa" AND "empresa"::text = ANY(%s) ORDER BY "empresa", "tipo"',
            (empresas,),
        )
        return [{**f, "banco": bancos.nombre(f["bancoCodigo"])} for f in cur.fetchall()]
